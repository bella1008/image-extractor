"""Read-only v1 reconciliation and reversible, non-active v2 checklist drafts.

This module does not evaluate PDF wording or approve migrated rules.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

SOURCE_HEADERS = (
    "common_id", "check_id", "topic", "scope", "exclude_scope", "source_token",
    "language", "doc_type", "section_heading", "block_type", "expected_result",
    "required_text", "match_method", "importance", "status", "evidence_file", "note",
)
FIELD_MAP = {"status": "approval_status", "source_token": "source_reference_token",
             "section_heading": "legacy_section_heading", "block_type": "legacy_block_type"}
DRAFT_HEADERS = tuple(FIELD_MAP.get(name, name) for name in SOURCE_HEADERS)
SCOPE_ALIASES = {name: name for name in (
    "WW", "ZC", "ZA", "ZX", "ASIA", "XY", "LATIN", "XC", "XH", "XU", "ZG", "ZH", "KR", "EU",
)} | {"GLOBAL": "WW"}


class MigrationError(ValueError):
    """Input or draft cannot be migrated without losing or changing source data."""


def read_excel_rows(path: Path, *, draft: bool = False) -> list[dict[str, str]]:
    from openpyxl import load_workbook
    headers = DRAFT_HEADERS if draft else SOURCE_HEADERS
    sheet_name = "Checklist_V2" if draft else "Checklist_Master"
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise MigrationError(f"missing sheet: {sheet_name}")
        sheet = workbook[sheet_name]
        rows = iter(sheet.iter_rows())
        first = next(rows, ())
        actual = tuple(cell.value for cell in first)
        if actual != headers:
            raise MigrationError(f"unexpected columns/order in {sheet_name}: {actual}")
        result = []
        for excel_row, cells in enumerate(rows, start=2):
            values = []
            for cell in cells:
                if cell.data_type in ("f", "e"):
                    raise MigrationError(f"formula/error cell: {sheet_name}!{cell.coordinate}")
                if cell.value is not None and not isinstance(cell.value, str):
                    raise MigrationError(f"non-text rule cell: {sheet_name}!{cell.coordinate}")
                values.append(cell.value if cell.value is not None else "")
            if not any(values):
                raise MigrationError(f"blank row interrupts master at Excel row {excel_row}")
            result.append(dict(zip(headers, values, strict=True)))
        validate_rows(restore_source_rows(result) if draft else result)
        return result
    finally:
        workbook.close()


def validate_rows(rows: list[dict[str, str]]) -> None:
    if not rows:
        raise MigrationError("empty checklist master")
    seen = set()
    for excel_row, row in enumerate(rows, 2):
        if set(row) != set(SOURCE_HEADERS) or any(not isinstance(v, str) for v in row.values()):
            raise MigrationError(f"unexpected fields/types at row {excel_row}")
        for field in ("common_id", "check_id", "scope", "language", "required_text"):
            if not row[field].strip():
                raise MigrationError(f"empty {field} at row {excel_row}")
        key = row["check_id"].strip()
        if key in seen:
            raise MigrationError(f"duplicate check_id: {key}")
        seen.add(key)
        if row["status"] not in ("approved", "review", "deprecated"):
            raise MigrationError(f"unsupported approval state: {key}")
        if row["match_method"] not in ("presence", "normalized_exact", "table_row", "table_header"):
            raise MigrationError(f"unsupported match_method: {key}")
        if row["expected_result"] != "required_present":
            raise MigrationError(f"unsupported expected_result: {key}")


def split_values(value: str) -> list[str]:
    return [part.strip() for part in value.split(";") if part.strip()]


def _legacy_scope(value: str) -> list[str]:
    result = []
    for part in split_values(value):
        key = part.upper()
        if len(part) <= 6 and key not in SCOPE_ALIASES and "_" not in part and " " not in part:
            raise MigrationError(f"unknown v1 scope alias: {part}")
        normalized = SCOPE_ALIASES.get(key, part)
        if normalized not in result:
            result.append(normalized)
    return result


def legacy_rule(row: dict[str, str]) -> dict[str, Any]:
    """Reproduce the frozen v1 exporter, only for reconciliation with existing JSON."""
    clean = {key: value.strip() for key, value in row.items()}
    return {**clean, "scope": _legacy_scope(clean["scope"]),
            "exclude_scope": _legacy_scope(clean["exclude_scope"]), "language": split_values(clean["language"])}


def sync_differences(rows: list[dict[str, str]], payload: dict) -> list[dict]:
    validate_rows(rows)
    if type(payload.get("schema_version")) is not int or payload["schema_version"] != 1:
        raise MigrationError("unsupported source JSON schema")
    actual = payload.get("rules")
    if not isinstance(actual, list) or any(not isinstance(r, dict) for r in actual):
        raise MigrationError("invalid source JSON rules")
    differences = []
    if len(rows) != len(actual):
        differences.append({"excel_row": None, "field": "row_count", "excel": len(rows), "json": len(actual)})
    for index, (source, runtime) in enumerate(zip(rows, actual), 2):
        expected = legacy_rule(source)
        for key in sorted(expected.keys() | runtime.keys()):
            if key not in expected or key not in runtime or expected.get(key) != runtime.get(key):
                differences.append({"excel_row": index, "check_id": source["check_id"], "field": key,
                                    "excel": expected.get(key), "json": runtime.get(key)})
    return differences


def draft_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    validate_rows(rows)
    return [{FIELD_MAP.get(key, key): row[key] for key in SOURCE_HEADERS} for row in rows]


def restore_source_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    result = []
    for row in rows:
        if set(row) != set(DRAFT_HEADERS):
            raise MigrationError("draft fields do not match reversible mapping")
        result.append({key: row[FIELD_MAP.get(key, key)] for key in SOURCE_HEADERS})
    return result


def metadata_applies(row: dict[str, str], context: dict[str, str], *, restrict_source_token: bool = False) -> bool:
    """Mirror v1 metadata_matches, with an explicit diagnostic-only counterfactual."""
    source_token, region = context["source_token"], context["region"]
    excluded = split_values(row["exclude_scope"])
    languages = split_values(row["language"])
    scopes = split_values(row["scope"])
    applies = (
        source_token not in excluded and region not in excluded
        and (not languages or context["language"] in languages)
        and (not row["doc_type"] or row["doc_type"] == context["doc_type"])
        and bool({"WW", "GLOBAL", "COMMON", source_token, region}.intersection(scopes))
    )
    if restrict_source_token and row["source_token"]:
        applies = applies and source_token in split_values(row["source_token"])
    return applies


def profile_audit(rows: list[dict[str, str]], profiles: list[dict]) -> dict:
    validate_rows(rows)
    differences, overlaps, profile_counts = [], [], []
    applicable_counts, narrowing_counts, overlap_counts = Counter(), Counter(), Counter()
    for profile in profiles:
        languages = split_values(profile["languages"])
        for language in languages:
            context = {"source_token": profile["source_token"], "region": profile["region"],
                       "doc_type": profile["doc_type"], "language": language}
            groups = defaultdict(list)
            applicable = 0
            narrower = 0
            for excel_row, row in enumerate(rows, 2):
                if row["status"] != "approved" or not metadata_applies(row, context):
                    continue
                applicable += 1
                key = row["check_id"]
                applicable_counts[key] += 1
                groups[row["common_id"]].append(key)
                if not metadata_applies(row, context, restrict_source_token=True):
                    narrowing_counts[key] += 1
                    narrower += 1
                    differences.append({**context, "excel_row": excel_row, "check_id": key,
                                        "common_id": row["common_id"], "rule_source_token": row["source_token"],
                                        "legacy_applicable": True, "if_source_token_restricts": False})
            for common_id, keys in groups.items():
                if len(keys) > 1:
                    overlaps.append({**context, "common_id": common_id, "check_ids": keys})
                    overlap_counts.update(keys)
            profile_counts.append({**context, "legacy_applicable_rules": applicable,
                                   "if_source_token_restricts": applicable - narrower})
    audit_rows = []
    for excel_row, row in enumerate(rows, 2):
        key = row["check_id"]
        audit_rows.append({"excel_row": excel_row, "common_id": row["common_id"], "check_id": key,
                           "approval_status": row["status"],
                           "migration_status": "pending" if row["status"] == "approved" else "excluded",
                           "legacy_block_type": row["block_type"], "legacy_section_heading": row["section_heading"],
                           "legacy_match_method": row["match_method"], "applicable_profile_languages": applicable_counts[key],
                           "source_restriction_would_drop": narrowing_counts[key], "overlap_contexts": overlap_counts[key],
                           "reason": "XML selector/matching contract not verified" if row["status"] == "approved" else "Original approval state retained; not active"})
    return {"rows": audit_rows, "scope_differences": differences, "overlaps": overlaps, "profiles": profile_counts}


def evidence_inventory(rows: list[dict[str, str]], project_root: Path) -> list[dict]:
    root = Path(project_root).resolve()
    result = []
    for row in rows:
        references = []
        for value in split_values(row["evidence_file"]):
            normalized = value.replace("\\", "/")
            if not normalized.startswith(("outputs/", "samples/", "docs/", "metadata/")) and not Path(value).is_absolute():
                state = "descriptive_reference"
            else:
                path = (root / value).resolve()
                state = "outside_project" if not path.is_relative_to(root) else "exists_not_revalidated" if path.is_file() else "directory_reference" if path.is_dir() else "missing"
            references.append({"reference": value, "state": state})
        states = {item["state"] for item in references}
        state = next((name for name in ("outside_project", "missing", "descriptive_reference", "directory_reference") if name in states), "exists_not_revalidated" if states else "empty_reference")
        result.append({"check_id": row["check_id"], "evidence_file": row["evidence_file"],
                       "evidence_state": state, "references": references})
    return result

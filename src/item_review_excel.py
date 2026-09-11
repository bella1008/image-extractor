"""Item-result display contract; source and business decisions stay unchanged."""
import json
import re
import os
from io import BytesIO
from pathlib import Path
import subprocess
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.item_review_service import (_validate_report, _require, _digest, _unchanged, _write_new,
                                     read_completed_item_review)


def _cell(value):
    if type(value) is int:
        return value
    if not isinstance(value, str) or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', value):
        raise ValueError('Excel cannot preserve this cell value')
    if len(value.encode('utf-16-le')) // 2 > 32767:
        raise ValueError('Excel cell exceeds 32767 characters; source was not truncated')
    return value


def build_item_excel_view(report: dict) -> dict:
    """Build three exact matrices. Evidence IDs are report-local, not XML IDs."""
    _validate_report(report)
    summary = [['검토 범위', '부모 체크항목', report['parent_check_id']],
               ['검토 범위', '상태', '모든 항목 검토 필요. 모델 적용 미확정. 자동 합격·불합격 아님.'],
               ['사용 방법', '검토 메모', '결과 사본의 메모이며 DB 승인 입력이 아닙니다. 원장 제안은 자동 실행하지 않습니다.']]
    for key, label in [('item_count', '하위 항목'), ('found', '단일 문구 근거'), ('ambiguous', '복수 근거'),
                       ('not_found', '검색 범위 내 미발견'), ('not_examined', '범위 미확정'),
                       ('needs_review', '검토 필요'), ('condition_candidate_items', '조건 안내 연결 후보')]:
        summary.append(['집계', label, report['summary'][key]])
    for section, source in [('검토 대상', report['target_source']), ('원장 작성 당시', report['master_source'])]:
        for key, label in [('pdf_filename', 'PDF 파일'), ('pdf_sha256', 'PDF SHA256'),
                           ('semantic_xml_sha256', 'XML SHA256'), ('bundle_path', '추출 결과 위치'),
                           ('receipt_ref', '추출 완료 기록'), ('receipt_sha256', '완료 기록 SHA256')]:
            if key in source:
                summary.append([section, label, source[key]])
    items, evidence = [], []
    for item in report['items']:
        references = []
        for field, kind in [('candidates', '항목 문구'), ('condition_candidates', '조건 안내 후보')]:
            ids = []
            for window in item[field]:
                eid = f'E{len(evidence) + 1:04}'
                ids.append(eid)
                entries = [e for p in window['parts'] for e in p.get('evidence', [])]
                pages = sorted({e['page_index'] + 1 for e in entries if type(e.get('page_index')) is int})
                paths = list(dict.fromkeys(e['xml_path'] for e in entries if e.get('xml_path')))
                evidence.append([eid, item['item_key'], kind, window['text'], '\n'.join(window['owner_ids']),
                    '\n'.join(window['structure_types']), ', '.join(map(str, pages)), '\n'.join(paths),
                    json.dumps(window, ensure_ascii=False, sort_keys=True)])
            references.append('\n'.join(ids))
        author = item['author_proposal']
        items.append([item['item_key'], item['required_text'], '검토 필요', item['description'],
                      *references, '', author['proposed_model_rule'], author['proposal_evidence'], author['reviewer_note']])
    sheets = [dict(name='Summary', headers=['구분', '항목', '값'], rows=summary, widths=[20, 28, 110], freeze='A2'),
              dict(name='Item Results', headers=['고정 항목 키', '기준 문구', '검토 판정', '설명', '항목 근거 ID',
                   '조건 후보 ID', '검토 메모', '원장 모델 조건 제안', '원장 제안 근거', '원장 검토 메모'],
                   rows=items, widths=[30, 58, 14, 65, 20, 20, 45, 55, 55, 45], freeze='B2'),
              dict(name='Source Evidence', headers=['근거 ID', '고정 항목 키', '근거 종류', '현재 원문', '현재 노드 ID',
                   '태그 종류', 'PDF 페이지', 'XML 경로', '상세 근거 JSON'], rows=evidence,
                   widths=[14, 30, 18, 85, 45, 20, 12, 70, 230], freeze='C2')]
    for sheet in sheets:
        sheet['rows'] = [[_cell(v) for v in row] for row in sheet['rows']]
    return {'schema_version': 'item-review-excel-view/1', 'activation_status': 'draft_only',
            'decision_status': 'not_evaluated', 'sheets': sheets}


def validate_item_workbook(content: bytes, view: dict) -> None:
    """Read-only whole-cell and native-layout verification after authoring."""
    with ZipFile(BytesIO(content)) as archive:
        _require(not any('externallinks/' in n.lower() or 'vbaproject' in n.lower() for n in archive.namelist()),
                 'Excel external links/macros are not allowed')
        for name in archive.namelist():
            if name.endswith('.rels'):
                import xml.etree.ElementTree as ET
                _require(not any(e.get('TargetMode') == 'External' for e in ET.fromstring(archive.read(name))),
                         'Excel external relationship is not allowed')
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        _require(workbook.sheetnames == [s['name'] for s in view['sheets']], 'Excel sheet contract differs')
        for sheet in view['sheets']:
            ws = workbook[sheet['name']]
            expected = [sheet['headers'], *sheet['rows']]
            _require(ws.max_row == len(expected) and ws.max_column == len(sheet['headers']), 'Excel dimensions differ')
            _require(ws.sheet_state == 'visible' and not ws.merged_cells.ranges, 'Excel hidden or merged source cells')
            _require(ws.freeze_panes == sheet['freeze'], 'Excel frozen panes differ')
            ref = f'A1:{get_column_letter(len(sheet["headers"]))}{len(expected)}'
            filters = [ws.auto_filter.ref, *(table.autoFilter.ref for table in ws.tables.values() if table.autoFilter)]
            _require(ref in filters, 'Excel filter differs')
            for actual, wanted in zip(ws.iter_rows(), expected, strict=True):
                for cell, value in zip(actual, wanted, strict=True):
                    _require(cell.data_type not in ('f', 'e') and cell.hyperlink is None, 'Excel formula/error/link is not allowed')
                    observed = '' if cell.value is None else cell.value
                    _require(type(observed) is type(value) and observed == value,
                             f'Excel value differs: {sheet["name"]}!{cell.coordinate}')
    finally:
        workbook.close()


RUN_FILES = ('item_review_complete.json', 'item_observation.json', 'item_review.html')
EXCEL_FILES = ('item_excel_complete.json', 'item_excel_view.json', 'item_review.xlsx')


def _source_snapshot(run_dir):
    run_dir = Path(run_dir)
    snapshot = {run_dir / name: (run_dir / name).read_bytes() for name in RUN_FILES}
    try:
        report = read_completed_item_review(run_dir)
    except (AttributeError, TypeError, KeyError) as exc:
        raise ValueError('invalid source observation structure') from exc
    _unchanged(snapshot)
    return report, snapshot


def _run_builder(view_path, output_dir, node_executable, node_modules):
    """Explicit local runtime only; no installer, shell or remote execution."""
    node_executable, node_modules = Path(node_executable).resolve(), Path(node_modules).resolve()
    _require(node_executable.is_file() and node_modules.is_dir(), 'Excel authoring runtime is not configured')
    script = Path(__file__).resolve().parents[1] / 'scripts/build_item_review_excel.mjs'
    result = subprocess.run([str(node_executable), str(script), str(view_path.resolve()), str(output_dir.resolve()),
                             _digest(view_path.read_bytes()), str(node_modules)],
                            capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=180)
    _write_new(output_dir / 'authoring.log', result.stdout + '\n' + result.stderr)
    _require(result.returncode == 0, 'Excel authoring failed; see authoring.log')


def export_item_review_excel(run_dir: Path, output_dir: Path, node_executable: Path, node_modules: Path) -> dict:
    output_dir, run_dir = Path(output_dir), Path(run_dir)
    output_dir.mkdir(parents=True, exist_ok=False)
    published = False
    final = output_dir / EXCEL_FILES[0]
    try:
        report, source = _source_snapshot(run_dir)
        view = build_item_excel_view(report)
        view_path = output_dir / EXCEL_FILES[1]
        view_text = json.dumps(view, ensure_ascii=True, indent=2) + '\n'
        _write_new(view_path, view_text)
        expected = {view_path: view_text.encode('utf-8')}
        _run_builder(view_path, output_dir, node_executable, node_modules)
        path = output_dir / EXCEL_FILES[2]
        content = path.read_bytes()
        validate_item_workbook(content, view)
        expected[path] = content
        receipt = {'schema_version': 'item-review-excel-run/1', 'status': 'ready_for_human_review',
                   'decision_status': 'not_evaluated', 'activation_status': 'draft_only',
                   'source_completion_sha256': _digest(source[run_dir / RUN_FILES[0]]),
                   'artifacts': {p.name: _digest(b) for p, b in expected.items()}}
        def verify():
            _unchanged(source)
            _unchanged(expected)
            _require(not (run_dir / 'item_review_failed.json').exists(), 'source run failed')
            _require(not (output_dir / 'item_excel_failed.json').exists(), 'Excel export failed')
        verify()
        pending = output_dir / 'item_excel_complete.pending.json'
        text = json.dumps(receipt, ensure_ascii=True, indent=2) + '\n'
        _write_new(pending, text)
        expected[pending] = text.encode('utf-8')
        verify()
        os.link(pending, final)
        published = True
        expected[final] = expected[pending]
        verify()
        pending.unlink()
        return receipt
    except Exception as exc:
        if published:
            try: final.unlink()
            except OSError: pass
        try:
            _write_new(output_dir / 'item_excel_failed.json', json.dumps({'status': 'failed', 'reason': str(exc)}))
        except OSError: pass
        raise


def read_completed_item_excel(output_dir: Path, run_dir: Path) -> bytes:
    """Download only a completed workbook tied to the supplied observation run."""
    output_dir, run_dir = Path(output_dir), Path(run_dir)
    _require(not (output_dir / 'item_excel_failed.json').exists(), 'Excel export failed')
    report, source = _source_snapshot(run_dir)
    snapshot = {output_dir / n: (output_dir / n).read_bytes() for n in EXCEL_FILES}
    receipt = json.loads(snapshot[output_dir / EXCEL_FILES[0]])
    _require(isinstance(receipt, dict), 'invalid Excel completion structure')
    _require(receipt.get('schema_version') == 'item-review-excel-run/1'
             and receipt.get('status') == 'ready_for_human_review'
             and receipt.get('decision_status') == 'not_evaluated'
             and receipt.get('activation_status') == 'draft_only', 'invalid Excel completion state')
    _require(receipt.get('source_completion_sha256') == _digest(source[run_dir / RUN_FILES[0]]),
             'Excel belongs to a different observation run')
    _require(isinstance(receipt.get('artifacts'), dict)
             and set(receipt['artifacts']) == set(EXCEL_FILES[1:]), 'invalid Excel artifact set')
    for name, digest in receipt['artifacts'].items():
        _require(_digest(snapshot[output_dir / name]) == digest, f'Excel artifact changed: {name}')
    view = build_item_excel_view(report)
    _require(json.loads(snapshot[output_dir / EXCEL_FILES[1]]) == view, 'Excel view differs from current observation')
    content = snapshot[output_dir / EXCEL_FILES[2]]
    validate_item_workbook(content, view)
    _unchanged(snapshot)
    _unchanged(source)
    _require(not (output_dir / 'item_excel_failed.json').exists()
             and not (run_dir / 'item_review_failed.json').exists(), 'run failed during read')
    return content

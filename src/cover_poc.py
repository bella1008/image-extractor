from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.cli import DEFAULT_MAPPING_PATH
from src.pdf_analyzer import analyze_pdf_structure
from src.profile_lookup import ProfileLookupService
from src.profile_repository import PdfProfileRepository
from src.text_extractor import extract_review_text


@dataclass(frozen=True)
class CoverPocPaths:
    root: Path
    json_path: Path
    markdown_path: Path
    xlsx_path: Path
    page_image_path: Path
    region_dir: Path


def make_output_paths(output_dir: Path, source_token: str, language: str) -> CoverPocPaths:
    root = output_dir / f"{source_token}_{language}"
    return CoverPocPaths(
        root=root,
        json_path=root / "cover_poc.json",
        markdown_path=root / "cover_poc.md",
        xlsx_path=root / "cover_regions.xlsx",
        page_image_path=root / "cover_page_001.png",
        region_dir=root / "regions",
    )


def safe_region_slug(region: dict[str, Any]) -> str:
    label = region.get("heading") or region.get("text") or region["region_type"]
    label = re.sub(r"[^A-Za-z0-9]+", "_", label).strip("_").lower()
    label = label[:40].strip("_") or "region"
    return f"region_{region['region_order']:02d}_{region['region_type']}_{label}"


def build_cover_poc_payload(extracted: dict[str, Any]) -> dict[str, Any]:
    if not extracted.get("cover_pages"):
        raise ValueError("No cover page found in extracted payload.")

    cover = extracted["cover_pages"][0]
    regions = []
    for region in cover.get("regions", []):
        manual_review_required = is_manual_review_region(region)
        regions.append(
            {
                **region,
                "manual_review_required": manual_review_required,
                "ocr_recommended": manual_review_required,
                "crop_path": None,
            }
        )

    return {
        "file_name": extracted["file_name"],
        "source_token": extracted["source_token"],
        "region": extracted["region"],
        "buyer_codes": extracted["buyer_codes"],
        "language": extracted["language"],
        "doc_type": extracted["doc_type"],
        "detected_doc_type": extracted["detected_doc_type"],
        "cover": {
            "page_number": cover["page_number"],
            "page_role": cover["page_role"],
            "schema": cover.get("schema", {}),
            "schema_validation": cover.get("schema_validation", {}),
            "markdown": cover.get("markdown", ""),
        },
        "summary": {
            "region_count": len(regions),
            "manual_review_region_count": sum(
                1 for region in regions if region["manual_review_required"]
            ),
            "table_region_count": sum(1 for region in regions if "table" in region),
        },
        "regions": regions,
    }


def is_manual_review_region(region: dict[str, Any]) -> bool:
    if region["region_type"] in {"contact_table", "model_serial_fields"}:
        return True
    if "table" in region:
        return True
    return region["region_type"] in {"body", "disclaimer", "support_note"}


def build_cover_review_markdown(payload: dict[str, Any]) -> str:
    cover = payload["cover"]
    validation = cover.get("schema_validation", {})
    lines = [
        "# Cover Page POC",
        "",
        f"- file_name: `{payload['file_name']}`",
        f"- source_token: `{payload['source_token']}`",
        f"- language: `{payload['language']}`",
        f"- doc_type: `{payload['doc_type']}`",
        f"- cover_page: `{cover['page_number']}`",
        f"- schema_id: `{cover.get('schema', {}).get('schema_id', '')}`",
        f"- schema_status: `{validation.get('status', '')}`",
        f"- missing_required_regions: `{';'.join(validation.get('missing_required_regions', []))}`",
        f"- region_count: `{payload['summary']['region_count']}`",
        f"- manual_review_region_count: `{payload['summary']['manual_review_region_count']}`",
        "",
        "## Region Review",
    ]
    for region in payload["regions"]:
        lines.extend(
            [
                "",
                f"### {region['region_order']:02d}. {region['region_type']}",
                "",
                f"- heading: `{region.get('heading') or ''}`",
                f"- bbox: `{','.join(str(value) for value in region['bbox'])}`",
                f"- crop_path: `{region.get('crop_path') or ''}`",
                f"- manual_review_required: `{str(region['manual_review_required']).lower()}`",
                f"- ocr_recommended: `{str(region['ocr_recommended']).lower()}`",
                "",
                region.get("markdown") or region["text"],
            ]
        )
    return "\n".join(lines) + "\n"


def write_cover_poc_outputs(
    payload: dict[str, Any], paths: CoverPocPaths
) -> dict[str, Path]:
    paths.root.mkdir(parents=True, exist_ok=True)
    paths.region_dir.mkdir(parents=True, exist_ok=True)

    paths.json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    paths.markdown_path.write_text(build_cover_review_markdown(payload), encoding="utf-8")
    write_cover_regions_xlsx(payload, paths.xlsx_path)

    return {
        "json_path": paths.json_path,
        "markdown_path": paths.markdown_path,
        "xlsx_path": paths.xlsx_path,
    }


def write_cover_regions_xlsx(payload: dict[str, Any], output_path: Path) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Cover Summary"
    write_cover_summary_sheet(summary_sheet, payload)
    write_cover_review_sheet(workbook.create_sheet("Cover Review"), payload)
    write_cover_markdown_sheet(workbook.create_sheet("Cover Markdown"), payload)
    write_compare_fields_sheet(workbook.create_sheet("Compare Fields"), payload)
    workbook.save(output_path)


def write_cover_summary_sheet(sheet, payload: dict[str, Any]) -> None:
    cover = payload["cover"]
    validation = cover.get("schema_validation", {})
    rows = [
        ("file_name", payload["file_name"]),
        ("language", payload["language"]),
        ("source_token", payload["source_token"]),
        ("doc_type", payload["doc_type"]),
        ("cover_page", cover["page_number"]),
        ("schema_id", cover.get("schema", {}).get("schema_id", "")),
        ("schema_status", validation.get("status", "")),
        ("missing_required_regions", ";".join(validation.get("missing_required_regions", []))),
        ("region_count", payload["summary"]["region_count"]),
        ("manual_review_region_count", payload["summary"]["manual_review_region_count"]),
        ("table_region_count", payload["summary"]["table_region_count"]),
    ]
    sheet.append(["field", "value"])
    for row in rows:
        sheet.append(list(row))
    style_sheet(sheet, widths=[28, 90])


def write_cover_review_sheet(sheet, payload: dict[str, Any]) -> None:
    headers = [
        "region_order",
        "region_type",
        "heading",
        "manual_review_required",
        "text",
        "crop_path",
    ]
    sheet.append(headers)
    for region in payload["regions"]:
        sheet.append(
            [
                region["region_order"],
                region["region_type"],
                region.get("heading") or "",
                region["manual_review_required"],
                region["text"],
                region.get("crop_path") or "",
            ]
        )
    style_sheet(sheet, widths=[13, 22, 32, 22, 120, 80])


def write_cover_markdown_sheet(sheet, payload: dict[str, Any]) -> None:
    headers = ["region_order", "region_type", "markdown"]
    sheet.append(headers)
    for region in payload["regions"]:
        sheet.append(
            [
                region["region_order"],
                region["region_type"],
                region.get("markdown") or "",
            ]
        )
    style_sheet(sheet, widths=[13, 22, 120])


def write_compare_fields_sheet(sheet, payload: dict[str, Any]) -> None:
    headers = ["region_order", "region_type", "field_schema", "field_name", "field_value", "label"]
    sheet.append(headers)
    for region in payload["regions"]:
        fields = region.get("normalized_fields")
        if not fields:
            continue
        field_schema = fields.get("field_schema", "")
        for field_name, value in fields.items():
            if field_name == "field_schema":
                continue
            for field_value, label in iter_compare_field_values(value):
                sheet.append(
                    [
                        region["region_order"],
                        region["region_type"],
                        field_schema,
                        field_name,
                        field_value,
                        label,
                    ]
                )
    style_sheet(sheet, widths=[13, 22, 28, 26, 80, 20])


def iter_compare_field_values(value: Any) -> list[tuple[str, str]]:
    if isinstance(value, list):
        rows = []
        for item in value:
            if isinstance(item, dict):
                rows.append((str(item.get("url", "")), str(item.get("label", ""))))
            else:
                rows.append((str(item), ""))
        return rows
    if isinstance(value, dict):
        return [(json.dumps(value, ensure_ascii=False), "")]
    if value is None:
        return []
    return [(str(value), "")]


def style_sheet(sheet, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def render_cover_images(pdf_path: Path, payload: dict[str, Any], paths: CoverPocPaths) -> None:
    paths.region_dir.mkdir(parents=True, exist_ok=True)
    page_number = payload["cover"]["page_number"]
    zoom = 2
    matrix = fitz.Matrix(zoom, zoom)
    with fitz.open(pdf_path) as document:
        page = document[page_number - 1]
        page.get_pixmap(matrix=matrix, alpha=False).save(paths.page_image_path)
        for region in payload["regions"]:
            rect = fitz.Rect(region["bbox"])
            clip = rect + (-3, -3, 3, 3)
            clip &= page.rect
            crop_path = paths.region_dir / f"{safe_region_slug(region)}.png"
            page.get_pixmap(matrix=matrix, clip=clip, alpha=False).save(crop_path)
            region["crop_path"] = str(crop_path)


def run_cover_poc(
    pdf_path: str | Path,
    output_dir: str | Path,
    language: str = "ENG",
    mapping_path: str | Path = DEFAULT_MAPPING_PATH,
) -> dict[str, Path]:
    pdf_path = Path(pdf_path)
    repository = PdfProfileRepository(mapping_path)
    lookup = ProfileLookupService(repository).lookup_for_file(str(pdf_path))
    structure = analyze_pdf_structure(str(pdf_path))
    extracted = extract_review_text(
        lookup.parsed_filename,
        lookup.profile,
        structure,
        language=language,
    )
    payload = build_cover_poc_payload(extracted)
    paths = make_output_paths(Path(output_dir), extracted["source_token"], language)
    render_cover_images(pdf_path, payload, paths)
    return write_cover_poc_outputs(payload, paths)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Generate a cover-page POC review package.")
    parser.add_argument("pdf_path", help="Path to the source PDF.")
    parser.add_argument("--language", default="ENG", help="Language code to extract.")
    parser.add_argument(
        "--output-dir",
        default="outputs/cover_poc",
        help="Root output directory. A source/language subfolder is created inside it.",
    )
    parser.add_argument(
        "--mapping",
        default=str(DEFAULT_MAPPING_PATH),
        help=f"Path to profile mapping JSON. Default: {DEFAULT_MAPPING_PATH}",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    written = run_cover_poc(
        args.pdf_path,
        args.output_dir,
        language=args.language,
        mapping_path=args.mapping,
    )
    for label, path in written.items():
        print(f"{label}: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

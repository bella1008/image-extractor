import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.content_poc import (
    build_content_markdown,
    build_content_poc_payload,
    make_output_paths,
    write_content_poc_outputs,
)


class ContentPocTests(unittest.TestCase):
    def test_make_output_paths_uses_source_token_and_language(self):
        paths = make_output_paths(Path("outputs/content_poc"), "ZC_L02", "ENG")

        self.assertEqual(paths.root, Path("outputs/content_poc/ZC_L02_ENG"))
        self.assertEqual(paths.json_path.name, "content_poc.json")
        self.assertEqual(paths.markdown_path.name, "content_blocks.md")
        self.assertEqual(paths.xlsx_path.name, "content_review.xlsx")
        self.assertEqual(paths.crop_dir.name, "crops")

    def test_build_content_payload_excludes_cover_cells_and_flags_risky_blocks(self):
        payload = build_content_poc_payload(minimal_extracted_payload())

        self.assertEqual(payload["summary"]["cell_count"], 2)
        self.assertEqual(payload["summary"]["block_count"], 3)
        self.assertEqual(payload["summary"]["manual_review_block_count"], 2)
        self.assertEqual(
            [block["block_type"] for block in payload["blocks"]],
            ["heading", "safety_symbol_table", "navigation_ui"],
        )
        self.assertTrue(payload["blocks"][1]["manual_review_required"])
        self.assertTrue(payload["blocks"][1]["image_crop_required"])
        self.assertEqual(payload["blocks"][1]["source"]["cell_role"], "content")

    def test_build_content_markdown_groups_blocks_by_section(self):
        payload = build_content_poc_payload(minimal_extracted_payload())

        markdown = build_content_markdown(payload)

        self.assertIn("# Content POC", markdown)
        self.assertIn("## 01. Warning! Important Safety Instructions", markdown)
        self.assertIn("### Block 002 - safety_symbol_table", markdown)
        self.assertIn("manual_review_required: `true`", markdown)

    def test_write_content_outputs_writes_review_workbook(self):
        payload = build_content_poc_payload(minimal_extracted_payload())

        with tempfile.TemporaryDirectory() as temp_dir:
            paths = make_output_paths(Path(temp_dir), "ZC_L02", "ENG")
            written = write_content_poc_outputs(payload, paths)

            self.assertTrue(written["json_path"].exists())
            self.assertTrue(written["markdown_path"].exists())
            self.assertTrue(written["xlsx_path"].exists())
            saved = json.loads(written["json_path"].read_text(encoding="utf-8"))
            self.assertEqual(saved["summary"]["block_count"], 3)

            workbook = load_workbook(written["xlsx_path"])
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Content Summary",
                    "Content Review",
                    "Safety Tables",
                    "Navigation UI",
                    "Spec Tables",
                ],
            )
            self.assertEqual(
                [cell.value for cell in workbook["Content Review"][1]],
                [
                    "block_order",
                    "section_heading",
                    "block_type",
                    "manual_review_required",
                    "text",
                    "sentence_count",
                    "crop_path",
                    "page",
                    "cell_order",
                ],
            )

    def test_embedded_ventilation_heading_is_split_from_previous_bullet(self):
        payload = build_content_poc_payload(minimal_extracted_payload_with_embedded_heading())

        headings = [block["text"] for block in payload["blocks"] if block["block_type"] == "heading"]
        self.assertIn("Providing proper ventilation for your TV", headings)
        ventilation_block = next(
            block
            for block in payload["blocks"]
            if block["section_heading"] == "Providing proper ventilation for your TV"
            and block["block_type"] == "bullet"
        )
        self.assertTrue(ventilation_block["text"].startswith("When installing the TV"))

    def test_caution_warning_row_is_part_of_safety_symbol_table(self):
        payload = build_content_poc_payload(minimal_extracted_payload_with_safety_warning_row())

        table_blocks = [
            block
            for block in payload["blocks"]
            if block["section_heading"] == "Warning! Important Safety Instructions"
            and block["block_type"] == "safety_symbol_table"
        ]
        self.assertEqual(len(table_blocks), 2)
        self.assertTrue(table_blocks[0]["text"].startswith("CAUTION RISK OF ELECTRIC SHOCK"))
        self.assertEqual(table_blocks[0]["rows"][0]["symbol_key"], "caution_warning")
        self.assertEqual(table_blocks[0]["rows"][0]["label"], "CAUTION")

    def test_additional_embedded_headings_are_split_from_body_text(self):
        payload = build_content_poc_payload(minimal_extracted_payload_with_additional_headings())

        headings = [block["text"] for block in payload["blocks"] if block["block_type"] == "heading"]

        self.assertIn("Preventing the TV from falling", headings)
        self.assertIn("Precautions when installing the TV with a stand", headings)
        self.assertIn("How to turn on and off the Microphone", headings)
        self.assertNotIn("(The Frame only)", headings)

    def test_specification_compound_heading_is_split_into_spec_table(self):
        payload = build_content_poc_payload(minimal_extracted_payload_with_specifications())

        headings = [block["text"] for block in payload["blocks"] if block["block_type"] == "heading"]
        spec_tables = [block for block in payload["blocks"] if block["block_type"] == "spec_table"]

        self.assertIn("04 Specifications and Other Information", headings)
        self.assertIn("Specifications", headings)
        self.assertEqual(len(spec_tables), 1)
        self.assertEqual(spec_tables[0]["section_heading"], "Specifications")
        self.assertEqual(spec_tables[0]["rows"][0]["field"], "Display Resolution")
        self.assertIn("QN9**H", spec_tables[0]["rows"][0]["value"])


def minimal_extracted_payload():
    return {
        "file_name": "sample.pdf",
        "source_token": "ZC_L02",
        "region": "ZC",
        "buyer_codes": ["ZC"],
        "language": "ENG",
        "doc_type": "A2",
        "detected_doc_type": "A2",
        "pages": [
            {
                "page_number": 1,
                "review_order": 1,
                "language": "ENG",
                "page_role": "content",
                "cells": [
                    {
                        "row": 1,
                        "column": 1,
                        "role": "cover",
                        "reading_order": 1,
                        "reading_direction": "LTR",
                        "review_items": [{"item_order": 1, "kind": "body", "text": "Cover", "sentences": ["Cover"]}],
                    },
                    {
                        "row": 1,
                        "column": 3,
                        "role": "content",
                        "reading_order": 3,
                        "reading_direction": "LTR",
                        "review_items": [
                            {
                                "item_order": 1,
                                "kind": "heading",
                                "text": "Warning! Important Safety Instructions",
                                "sentences": [],
                            },
                            {
                                "item_order": 2,
                                "kind": "body",
                                "text": (
                                    "This symbol indicates that high voltage is present inside. "
                                    "Class II product: This symbol indicates that a safety connection "
                                    "to electrical earth is not required. AC voltage: Rated voltage "
                                    "marked with this symbol is AC voltage. DC voltage: Rated voltage "
                                    "marked with this symbol is DC voltage. Caution. Consult instructions "
                                    "for use: This symbol instructs the user to consult the user manual."
                                ),
                                "sentences": ["This symbol indicates that high voltage is present inside."],
                            },
                        ],
                    },
                    {
                        "row": 1,
                        "column": 4,
                        "role": "content",
                        "reading_order": 4,
                        "reading_direction": "LTR",
                        "review_items": [
                            {
                                "item_order": 1,
                                "kind": "navigation",
                                "text": "> Settings > Support > Open User Guide",
                                "sentences": ["> Settings > Support > Open User Guide"],
                            }
                        ],
                    },
                ],
            }
        ],
    }


def minimal_extracted_payload_with_embedded_heading():
    payload = minimal_extracted_payload()
    payload["pages"][0]["cells"][1]["review_items"] = [
        {
            "item_order": 1,
            "kind": "heading",
            "text": "Mounting the TV on a wall",
            "sentences": [],
        },
        {
            "item_order": 2,
            "kind": "bullet",
            "text": (
                "For 82 inch or larger models, have four people mount the TV onto a wall. "
                "Providing proper ventilation for your TV When installing the TV, ensure proper ventilation."
            ),
            "sentences": [
                "For 82 inch or larger models, have four people mount the TV onto a wall.",
                "Providing proper ventilation for your TV When installing the TV, ensure proper ventilation.",
            ],
        },
    ]
    payload["pages"][0]["cells"] = payload["pages"][0]["cells"][:2]
    return payload


def minimal_extracted_payload_with_safety_warning_row():
    payload = minimal_extracted_payload()
    payload["pages"][0]["cells"][1]["review_items"] = [
        {
            "item_order": 1,
            "kind": "heading",
            "text": "Warning! Important Safety Instructions",
            "sentences": [],
        },
        {
            "item_order": 2,
            "kind": "warning",
            "text": "CAUTION RISK OF ELECTRIC SHOCK. DO NOT OPEN.",
            "sentences": ["CAUTION RISK OF ELECTRIC SHOCK.", "DO NOT OPEN."],
        },
        {
            "item_order": 3,
            "kind": "body",
            "text": (
                "This symbol indicates that high voltage is present inside. "
                "Class II product: This symbol indicates that a safety connection "
                "to electrical earth is not required. AC voltage: Rated voltage "
                "marked with this symbol is AC voltage. DC voltage: Rated voltage "
                "marked with this symbol is DC voltage. Caution. Consult instructions "
                "for use: This symbol instructs the user to consult the user manual."
            ),
            "sentences": ["This symbol indicates that high voltage is present inside."],
        },
    ]
    payload["pages"][0]["cells"] = payload["pages"][0]["cells"][:2]
    return payload


def minimal_extracted_payload_with_additional_headings():
    payload = minimal_extracted_payload()
    payload["pages"][0]["cells"][1]["review_items"] = [
        {
            "item_order": 1,
            "kind": "heading",
            "text": "WARNING",
            "sentences": [],
        },
        {
            "item_order": 2,
            "kind": "body",
            "text": (
                "Preventing the TV from falling : Wall-anchor (not supplied) "
                "Precautions when installing the TV with a (The Frame only) "
                "When you install the TV with a stand, avoid placing the stand on the back part."
            ),
            "sentences": [
                "Preventing the TV from falling : Wall-anchor (not supplied) Precautions when installing the TV with a stand (The Frame only) When you install the TV with a stand, avoid placing the stand on the back part."
            ],
        },
        {
            "item_order": 3,
            "kind": "body",
            "text": "How to turn on and off the Microphone Type A Type B",
            "sentences": ["How to turn on and off the Microphone Type A Type B"],
        },
    ]
    payload["pages"][0]["cells"] = payload["pages"][0]["cells"][:2]
    return payload


def minimal_extracted_payload_with_specifications():
    payload = minimal_extracted_payload()
    payload["pages"][0]["cells"][1]["review_items"] = [
        {
            "item_order": 1,
            "kind": "heading",
            "text": "04 Specifications and Other Information Specifications Display Resolution",
            "sentences": [],
        },
        {
            "item_order": 2,
            "kind": "body",
            "text": "QN9**H: 7680 x 4320 Other models: 3840 x 2160",
            "sentences": ["QN9**H: 7680 x 4320 Other models: 3840 x 2160"],
        },
    ]
    payload["pages"][0]["cells"] = payload["pages"][0]["cells"][:2]
    return payload


if __name__ == "__main__":
    unittest.main()

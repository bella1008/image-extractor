"""The shared executable boundary; real authoring requires explicit bundled paths."""
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess

import pytest

ROOT = Path(__file__).resolve().parents[1]
NODE = os.environ.get('ITEM_REVIEW_NODE') or shutil.which('node')
BACKEND = ROOT / 'scripts/review_workbook_backend.mjs'
WRAPPER = ROOT / 'scripts/build_combined_review_excel.mjs'


def view():
    item = ['고정 항목 키', '기준 문구', '현재 원문', '검토 판정', '설명',
            'PDF 페이지', '조건 안내 원문 (후보)', '검토 메모']
    specs = [
        ('Summary', ['구분', '항목', '값'], [['범위', f'항목{i}', i] for i in range(11)], [20, 28, 110], 'A2'),
        ('Checklist Results', ['체크 ID', '기준 제목', '언어', '기준 문구', '현재 원문', '검토 판정', '설명', 'PDF 페이지', '검토 메모'],
         [['CHK-002', '구성품', 'ENG', '=1+1', "'literal", '검토 필요', '+cmd', '1', '@SUM(A1)']],
         [28, 30, 12, 58, 70, 14, 65, 12, 45], 'B2'),
        ('Item Results', item, [['key', '=1+1', '=1+1', '검토 필요', '+cmd', '1', '@SUM(A1)', "'literal"]],
         [30, 58, 58, 14, 65, 12, 65, 45], 'B2'),
        ('Source Evidence', ['체크 ID', '고정 항목 키', '근거 종류', '현재 원문', '현재 노드 ID', '태그 종류', 'PDF 페이지', 'XML 경로'],
         [['CHK-002', 'key', '구성품 문구', '=1+1', 'node1', 'paragraph', '1', '/p']],
         [28, 30, 18, 85, 45, 20, 12, 70], 'C2')]
    return dict(schema_version='combined-review-excel-view/2', activation_status='draft_only',
                decision_status='not_evaluated', sheets=[dict(name=n, headers=h, rows=r, widths=w, freeze=f)
                                                       for n, h, r, w, f in specs])


def run_node(*args):
    assert NODE, 'Node is required for backend boundary tests'
    return subprocess.run([NODE, *map(str, args)], cwd=ROOT, capture_output=True, text=True, encoding='utf-8')


def validate(data):
    assert BACKEND.is_file(), 'Reusable workbook backend missing'
    assert WRAPPER.is_file(), 'Combined trusted wrapper missing'
    code = (f'import {{validateView}} from {json.dumps(BACKEND.as_uri())};'
            f'import {{contract}} from {json.dumps(WRAPPER.as_uri())};'
            f'validateView({json.dumps(data)},contract);')
    return run_node('--input-type=module', '-e', code)


def test_shared_contract_accepts_valid_view_without_loading_authoring_runtime():
    result = validate(view())
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize('mutation, message', [
    ('schema', 'Unsupported Excel view'), ('decision', 'Unsupported Excel view'),
    ('sheets', 'Unexpected sheets'), ('headers', 'Unexpected headers'),
    ('height', 'Source needs more display space'), ('width', 'Invalid Excel widths'),
    ('freeze', 'Invalid freeze pane'), ('summary', 'Unexpected Summary rows')])
def test_shared_contract_fails_closed_before_authoring(mutation, message):
    data = view()
    if mutation == 'schema': data['schema_version'] = 'other/1'
    elif mutation == 'decision': data['decision_status'] = 'pass'
    elif mutation == 'sheets': data['sheets'].reverse()
    elif mutation == 'headers': data['sheets'][1]['headers'][0] = 'Other'
    elif mutation == 'height': data['sheets'][1]['rows'][0][3] = 'line\n' * 32
    elif mutation == 'width': data['sheets'][1]['widths'][0] = 0
    elif mutation == 'freeze': data['sheets'][1]['freeze'] = 'invalid'
    else: data['sheets'][0]['rows'].append(['unexpected', 'row', 1])
    result = validate(data)
    assert result.returncode != 0
    assert message in result.stderr


@pytest.mark.parametrize('wrapper', ['build_item_review_excel.mjs', 'build_combined_review_excel.mjs'])
def test_wrappers_reject_bad_hash_before_resolving_dependencies(tmp_path, wrapper):
    executable = ROOT / 'scripts' / wrapper
    assert executable.is_file(), 'Trusted wrapper missing'
    source = tmp_path / 'view.json'
    source.write_text('{}')
    result = run_node(executable, source, tmp_path, '0' * 64, tmp_path / 'missing_modules')
    assert result.returncode != 0
    assert 'Display contract hash differs' in result.stderr


@pytest.mark.parametrize('empty', [False, True], ids=['literal-text', 'empty-evidence'])
def test_real_combined_backend_preserves_cells_panes_filters_and_output(tmp_path, empty):
    node = os.environ.get('ITEM_REVIEW_NODE')
    modules = os.environ.get('ITEM_REVIEW_NODE_MODULES')
    if not node or not modules:
        pytest.skip('explicit bundled runtime and operation marker required for real authoring')
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    data = view()
    if empty: data['sheets'][-1]['rows'] = []
    source = tmp_path / 'view.json'
    content = json.dumps(data, ensure_ascii=False).encode()
    source.write_bytes(content)
    result = subprocess.run([node, str(WRAPPER), str(source), str(tmp_path),
                             hashlib.sha256(content).hexdigest(), modules],
                            capture_output=True, text=True, encoding='utf-8')
    assert result.returncode == 0, result.stderr
    output = tmp_path / 'review_report.xlsx'
    original = output.read_bytes()
    wb = load_workbook(output)
    assert wb.sheetnames == [s['name'] for s in data['sheets']]
    for ws, spec in zip(wb.worksheets, data['sheets'], strict=True):
        matrix = [spec['headers'], *spec['rows']]
        assert (ws.max_row, ws.max_column) == (len(matrix), len(spec['headers']))
        assert ws.freeze_panes == spec['freeze']
        ref = f'A1:{get_column_letter(len(spec["headers"]))}{len(matrix)}'
        assert any(t.autoFilter.ref == ref for t in ws.tables.values())
        for cells, values in zip(ws, matrix, strict=True):
            for cell, expected in zip(cells, values, strict=True):
                assert cell.data_type != 'f'
                assert (cell.value if cell.value is not None else '') == expected
        for row in ws.row_dimensions.values():
            assert row.height is None or row.height <= 409
    assert {p.name for p in tmp_path.glob('*.png')} >= {'summary.png', 'checklist.png', 'items.png', 'evidence.png'}
    retry = subprocess.run([node, str(WRAPPER), str(source), str(tmp_path),
                            hashlib.sha256(content).hexdigest(), modules], capture_output=True)
    assert retry.returncode != 0
    assert output.read_bytes() == original
    assert not list(tmp_path.glob('*.pending.xlsx'))

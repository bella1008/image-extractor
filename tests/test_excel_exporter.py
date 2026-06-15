import unittest

from openpyxl import Workbook

from src.excel_exporter import write_cover_regions


class ExcelExporterTests(unittest.TestCase):
    def test_cover_regions_sheet_includes_schema_validation_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        data = {
            "cover_pages": [
                {
                    "page_number": 1,
                    "schema": {"schema_id": "ZC_A2_COVER"},
                    "schema_validation": {
                        "status": "PASS",
                        "missing_required_regions": [],
                    },
                    "regions": [
                        {
                            "region_order": 1,
                            "region_type": "title",
                            "heading": "Simple User Guide",
                            "source": {"cell_order": 1},
                            "bbox": [1, 2, 3, 4],
                            "style_summary": {
                                "bold_candidate": True,
                                "max_font_size": 24,
                            },
                            "text": "Simple User Guide",
                            "markdown": "# Simple User Guide",
                        }
                    ],
                }
            ]
        }

        write_cover_regions(sheet, data)

        headers = [cell.value for cell in sheet[1]]
        self.assertIn("schema_id", headers)
        self.assertIn("schema_status", headers)
        self.assertIn("missing_required_regions", headers)
        self.assertEqual(sheet.cell(row=2, column=headers.index("schema_id") + 1).value, "ZC_A2_COVER")
        self.assertEqual(sheet.cell(row=2, column=headers.index("schema_status") + 1).value, "PASS")


if __name__ == "__main__":
    unittest.main()

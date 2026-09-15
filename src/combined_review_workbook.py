"""Python-only authoring of the combined reviewer workbook; no business decisions."""
from io import BytesIO
import math
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.combined_review_model import digest, parse_json, require
from src.item_review_excel import _cell, validate_item_workbook

SHEETS = (
    ('Summary', ['구분', '항목', '값']),
    ('Checklist Results', ['체크 ID', '기준 제목', '언어', '기준 문구', '현재 원문',
                           '검토 판정', '설명', 'PDF 페이지', '검토 메모']),
    ('Item Results', ['고정 항목 키', '기준 문구', '현재 원문', '검토 판정', '설명',
                      'PDF 페이지', '조건 안내 원문 (후보)', '검토 메모']),
    ('Source Evidence', ['체크 ID', '고정 항목 키', '근거 종류', '현재 원문', '현재 노드 ID',
                         '태그 종류', 'PDF 페이지', 'XML 경로']),
)


def _row_height(row, widths):
    def lines(value, width):
        return sum(max(1, math.ceil(sum(1.9 if ord(c) > 255 else 1 for c in line) / (width * .82)))
                   for line in str(value).split('\n'))
    return max(30, max(lines(v, w) for v, w in zip(row, widths, strict=True)) * 13 + 10)


def validate_view(view):
    require(isinstance(view, dict) and view.get('schema_version') == 'combined-review-excel-view/2'
            and view.get('activation_status') == 'draft_only'
            and view.get('decision_status') == 'not_evaluated', 'Unsupported Excel view')
    sheets = view.get('sheets')
    require(isinstance(sheets, list) and len(sheets) == len(SHEETS), 'Unexpected sheets')
    for spec, (name, headers) in zip(sheets, SHEETS, strict=True):
        require(isinstance(spec, dict) and spec.get('name') == name, 'Unexpected sheets')
        require(spec.get('headers') == headers, f'Unexpected headers: {name}')
        rows, widths = spec.get('rows'), spec.get('widths')
        require(isinstance(rows, list), 'Invalid Excel rows')
        require(name != 'Summary' or len(rows) == 11, 'Unexpected Summary rows')
        require(isinstance(widths, list) and len(widths) == len(headers)
                and all(type(w) in (int, float) and math.isfinite(w) and 0 < w <= 255 for w in widths),
                'Invalid Excel widths')
        require(spec.get('freeze') in ('A2', 'B2', 'C2'), 'Invalid freeze pane')
        for r, row in enumerate(rows, 2):
            require(isinstance(row, list) and len(row) == len(headers), 'Invalid Excel row')
            for value in row:
                _cell(value)
                require(type(value) is not int or abs(value) <= 2**53 - 1, 'Invalid Excel integer')
            require(_row_height(row, widths) <= 409, f'Source needs more display space: {name} row {r}')


def write_combined_workbook(view_path, output_dir, expected_hash):
    """Publish a checked XLSX without overwriting an earlier output.

    Only current combined view v2 is authored. Historical v1 workbooks remain
    readable through the existing archive reader, with their original bytes.
    """
    view_path, output_dir = Path(view_path), Path(output_dir)
    content = view_path.read_bytes()
    require(digest(content) == expected_hash, 'Display contract hash differs')
    view = parse_json(content)
    validate_view(view)
    output = output_dir / 'review_report.xlsx'
    if output.exists():
        raise FileExistsError(f'Workbook already exists: {output.name}')
    wb = Workbook()
    wb.remove(wb.active)
    try:
        for index, spec in enumerate(view['sheets'], 1):
            ws = wb.create_sheet(spec['name'])
            ws.sheet_view.showGridLines = False
            ws.freeze_panes = spec['freeze']
            for c, width in enumerate(spec['widths'], 1):
                ws.column_dimensions[get_column_letter(c)].width = width
            for r, values in enumerate([spec['headers'], *spec['rows']], 1):
                ws.row_dimensions[r].height = 32 if r == 1 else _row_height(values, spec['widths'])
                for c, value in enumerate(values, 1):
                    cell = ws.cell(r, c, value)
                    # Explicit string type preserves '=' and '#N/A' as source
                    # text, without a visible apostrophe or executable formula.
                    if isinstance(value, str):
                        cell.data_type = 's'
                    else:
                        cell.number_format = '#,##0'
                    cell.font = Font(name='Arial', size=10, color='202B3A', bold=r == 1)
                    cell.alignment = Alignment(wrap_text=True, vertical='top',
                                               horizontal='center' if r == 1 else None)
                    if r == 1:
                        cell.fill = PatternFill('solid', fgColor='D9EAF7')
                    elif spec['headers'][c - 1] in ('검토 판정', '검토 메모'):
                        cell.fill = PatternFill('solid', fgColor='FFF4CE')
            ref = f'A1:{get_column_letter(len(spec["headers"]))}{len(spec["rows"]) + 1}'
            table = Table(displayName=f'CombinedReviewTable{index}', ref=ref)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(table)
            if '검토 판정' in spec['headers']:
                ws.sheet_properties.tabColor = '243E60'
        stream = BytesIO()
        wb.save(stream)
        data = stream.getvalue()
        validate_item_workbook(data, view)
        require(view_path.read_bytes() == content, 'Display contract changed before publication')
        descriptor, name = tempfile.mkstemp(prefix='.combined-', suffix='.pending.xlsx', dir=output_dir)
        pending = Path(name)
        try:
            with os.fdopen(descriptor, 'wb') as target:
                target.write(data)
            require(view_path.read_bytes() == content, 'Display contract changed before publication')
            os.link(pending, output)
        finally:
            pending.unlink(missing_ok=True)
    finally:
        wb.close()

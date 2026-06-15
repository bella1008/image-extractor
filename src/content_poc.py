from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.cli import DEFAULT_MAPPING_PATH
from src.pdf_analyzer import analyze_pdf_structure
from src.profile_lookup import ProfileLookupService
from src.profile_repository import PdfProfileRepository
from src.text_extractor import (
    detect_block_type,
    extract_review_text,
    is_safety_symbol_table_text,
    parse_safety_symbol_rows,
)


EMBEDDED_HEADING_PATTERNS = [
    (
        "Providing proper ventilation for your TV",
        "Providing proper ventilation for your TV",
    ),
    ("Preventing the TV from falling", "Preventing the TV from falling"),
    (
        "Precautions when installing the TV with a",
        "Precautions when installing the TV with a stand",
    ),
    (
        "How to turn on and off the Microphone",
        "How to turn on and off the Microphone",
    ),
]

SPEC_COMPOUND_HEADING_PREFIX = (
    "04 Specifications and Other Information Specifications Display Resolution"
)


@dataclass(frozen=True)
class ContentPocPaths:
    root: Path
    json_path: Path
    markdown_path: Path
    xlsx_path: Path
    crop_dir: Path


def make_output_paths(output_dir: Path, source_token: str, language: str) -> ContentPocPaths:
    root = output_dir / f"{source_token}_{language}"
    return ContentPocPaths(
        root=root,
        json_path=root / "content_poc.json",
        markdown_path=root / "content_blocks.md",
        xlsx_path=root / "content_review.xlsx",
        crop_dir=root / "crops",
    )


def build_content_poc_payload(extracted: dict[str, Any]) -> dict[str, Any]:
    cells = collect_content_cells(extracted)
    blocks = build_content_blocks(cells)
    sections = build_content_sections(blocks)
    return {
        "file_name": extracted["file_name"],
        "source_token": extracted["source_token"],
        "region": extracted["region"],
        "buyer_codes": extracted["buyer_codes"],
        "language": extracted["language"],
        "doc_type": extracted["doc_type"],
        "detected_doc_type": extracted["detected_doc_type"],
        "summary": {
            "cell_count": len(cells),
            "block_count": len(blocks),
            "section_count": len(sections),
            "manual_review_block_count": sum(
                1 for block in blocks if block["manual_review_required"]
            ),
            "safety_table_count": sum(
                1 for block in blocks if block["block_type"] == "safety_symbol_table"
            ),
            "navigation_ui_count": sum(
                1 for block in blocks if block["block_type"] == "navigation_ui"
            ),
            "spec_table_count": sum(
                1 for block in blocks if block["block_type"] == "spec_table"
            ),
        },
        "sections": sections,
        "blocks": blocks,
    }


def collect_content_cells(extracted: dict[str, Any]) -> list[dict[str, Any]]:
    cells = []
    for page in extracted.get("pages", []):
        for cell in page.get("cells", []):
            if cell["role"] == "cover":
                continue
            cells.append(
                {
                    "page_number": page["page_number"],
                    "review_order": page["review_order"],
                    "page_role": page["page_role"],
                    "row": cell["row"],
                    "column": cell["column"],
                    "cell_role": cell["role"],
                    "cell_order": cell["reading_order"],
                    "reading_direction": cell["reading_direction"],
                    "review_items": cell.get("review_items", []),
                    "crop_path": None,
                }
            )
    return cells


def build_content_blocks(cells: list[dict[str, Any]]) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    current_section = "UNSECTIONED"
    section_order = 0
    for cell in cells:
        for item in expand_content_review_items(cell["review_items"]):
            for split_item in split_embedded_heading_items(item):
                block_type = detect_content_block_type(split_item, current_section)
                if block_type == "heading":
                    section_order += 1
                    current_section = split_item["text"]
                block = make_content_block(
                    split_item,
                    block_type,
                    cell,
                    block_order=len(blocks) + 1,
                    section_order=max(section_order, 1),
                    section_heading=current_section,
                )
                blocks.append(block)
    return blocks


def expand_content_review_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    expanded: list[dict[str, Any]] = []
    index = 0
    while index < len(items):
        item = items[index]
        next_item = items[index + 1] if index + 1 < len(items) else None
        spec_items = split_specification_heading_item(item, next_item)
        if spec_items:
            expanded.extend(spec_items)
            index += 2 if next_item else 1
            continue
        expanded.append(item)
        index += 1
    return expanded


def split_specification_heading_item(
    item: dict[str, Any], next_item: dict[str, Any] | None
) -> list[dict[str, Any]]:
    if item["kind"] != "heading" or item["text"] != SPEC_COMPOUND_HEADING_PREFIX:
        return []
    value = next_item["text"] if next_item and next_item["kind"] != "heading" else ""
    table_text = normalize_joined_text(f"Display Resolution {value}")
    return [
        {
            "kind": "heading",
            "text": "04 Specifications and Other Information",
            "sentences": [],
        },
        {"kind": "heading", "text": "Specifications", "sentences": []},
        {"kind": "spec_table", "text": table_text, "sentences": [table_text]},
    ]


def split_embedded_heading_items(item: dict[str, Any]) -> list[dict[str, Any]]:
    if item["kind"] in {"heading", "spec_table"}:
        return [item]
    remaining = [item]
    result: list[dict[str, Any]] = []
    while remaining:
        current = remaining.pop(0)
        if current["kind"] in {"heading", "spec_table"}:
            result.append(current)
            continue
        split_items = split_first_embedded_heading_item(current)
        if len(split_items) == 1 and split_items[0] is current:
            result.append(current)
        else:
            remaining = split_items + remaining
    return result


def split_first_embedded_heading_item(item: dict[str, Any]) -> list[dict[str, Any]]:
    text = item["text"]
    earliest: tuple[int, str, str] | None = None
    for trigger, heading in EMBEDDED_HEADING_PATTERNS:
        index = text.find(trigger)
        if index < 0:
            continue
        if earliest is None or index < earliest[0]:
            earliest = (index, trigger, heading)
    if earliest is None:
        return [item]

    index, trigger, heading = earliest
    before = text[:index].strip()
    after = text[index + len(trigger):].strip()
    split_items = []
    if before:
        split_items.append({**item, "text": before, "sentences": [before]})
    split_items.append({"kind": "heading", "text": heading, "sentences": []})
    if after:
        split_items.append({**item, "text": after, "sentences": [after]})
    return split_items


def normalize_joined_text(text: str) -> str:
    return " ".join(text.split())


def detect_content_block_type(item: dict[str, Any], section_heading: str) -> str:
    if item["kind"] == "spec_table":
        return "spec_table"
    if (
        item["kind"] == "warning"
        and section_heading == "Warning! Important Safety Instructions"
        and "CAUTION" in item["text"]
        and "RISK OF ELECTRIC SHOCK" in item["text"]
    ):
        return "safety_symbol_table"
    block_type = detect_block_type(item)
    if block_type == "body" and is_safety_symbol_table_text(item["text"]):
        return "safety_symbol_table"
    return block_type


def make_content_block(
    item: dict[str, Any],
    block_type: str,
    cell: dict[str, Any],
    block_order: int,
    section_order: int,
    section_heading: str,
) -> dict[str, Any]:
    risky = block_type in {"safety_symbol_table", "navigation_ui", "spec_table"}
    block = {
        "block_order": block_order,
        "section_order": section_order,
        "section_heading": section_heading,
        "block_type": block_type,
        "text": item["text"],
        "sentences": [
            {"sentence_order": index, "text": sentence}
            for index, sentence in enumerate(item.get("sentences", []), start=1)
        ],
        "manual_review_required": risky,
        "image_crop_required": risky,
        "ocr_recommended": risky,
        "crop_path": None,
        "source": {
            "page_number": cell["page_number"],
            "review_order": cell["review_order"],
            "page_role": cell["page_role"],
            "row": cell["row"],
            "column": cell["column"],
            "cell_role": cell["cell_role"],
            "cell_order": cell["cell_order"],
            "reading_direction": cell["reading_direction"],
        },
    }
    if block_type == "safety_symbol_table":
        block["table_type"] = "safety_symbol_table"
        rows = parse_safety_symbol_rows(item["text"])
        if not rows and "CAUTION" in item["text"]:
            rows = [
                {
                    "row_order": 0,
                    "symbol_key": "caution_warning",
                    "label": "CAUTION",
                    "description": item["text"],
                    "symbol_image_crop": None,
                }
            ]
        block["rows"] = rows
    if block_type == "spec_table":
        block["table_type"] = "spec_table"
        block["rows"] = parse_spec_table_rows(item["text"])
    return block


def parse_spec_table_rows(text: str) -> list[dict[str, str | int]]:
    normalized = normalize_joined_text(text)
    if normalized.startswith("Display Resolution "):
        return [
            {
                "row_order": 1,
                "field": "Display Resolution",
                "value": normalized.removeprefix("Display Resolution ").strip(),
            }
        ]
    return [{"row_order": 1, "field": "", "value": normalized}]


def build_content_sections(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    sections: dict[tuple[int, str], list[dict[str, Any]]] = {}
    for block in blocks:
        key = (block["section_order"], block["section_heading"])
        sections.setdefault(key, []).append(block)
    return [
        {
            "section_order": section_order,
            "heading": heading,
            "block_count": len(section_blocks),
            "manual_review_block_count": sum(
                1 for block in section_blocks if block["manual_review_required"]
            ),
        }
        for (section_order, heading), section_blocks in sections.items()
    ]


def build_content_markdown(payload: dict[str, Any]) -> str:
    lines = [
        "# Content POC",
        "",
        f"- file_name: `{payload['file_name']}`",
        f"- source_token: `{payload['source_token']}`",
        f"- language: `{payload['language']}`",
        f"- doc_type: `{payload['doc_type']}`",
        f"- block_count: `{payload['summary']['block_count']}`",
        f"- manual_review_block_count: `{payload['summary']['manual_review_block_count']}`",
    ]
    current_section = None
    for block in payload["blocks"]:
        if current_section != block["section_heading"]:
            current_section = block["section_heading"]
            lines.extend(["", f"## {block['section_order']:02d}. {current_section}"])
        lines.extend(
            [
                "",
                f"### Block {block['block_order']:03d} - {block['block_type']}",
                "",
                f"- page: `{block['source']['page_number']}`",
                f"- cell_order: `{block['source']['cell_order']}`",
                f"- crop_path: `{block.get('crop_path') or ''}`",
                f"- manual_review_required: `{str(block['manual_review_required']).lower()}`",
                "",
                block["text"],
            ]
        )
    return "\n".join(lines) + "\n"


def write_content_poc_outputs(
    payload: dict[str, Any], paths: ContentPocPaths
) -> dict[str, Path]:
    paths.root.mkdir(parents=True, exist_ok=True)
    paths.crop_dir.mkdir(parents=True, exist_ok=True)
    paths.json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    paths.markdown_path.write_text(build_content_markdown(payload), encoding="utf-8")
    write_content_review_xlsx(payload, paths.xlsx_path)
    return {
        "json_path": paths.json_path,
        "markdown_path": paths.markdown_path,
        "xlsx_path": paths.xlsx_path,
    }


def write_content_review_xlsx(payload: dict[str, Any], output_path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Content Summary"
    write_summary_sheet(summary, payload)
    write_content_review_sheet(workbook.create_sheet("Content Review"), payload)
    write_safety_tables_sheet(workbook.create_sheet("Safety Tables"), payload)
    write_navigation_ui_sheet(workbook.create_sheet("Navigation UI"), payload)
    write_spec_tables_sheet(workbook.create_sheet("Spec Tables"), payload)
    workbook.save(output_path)


def write_summary_sheet(sheet, payload: dict[str, Any]) -> None:
    sheet.append(["field", "value"])
    rows = [
        ("file_name", payload["file_name"]),
        ("language", payload["language"]),
        ("source_token", payload["source_token"]),
        ("doc_type", payload["doc_type"]),
        ("cell_count", payload["summary"]["cell_count"]),
        ("section_count", payload["summary"]["section_count"]),
        ("block_count", payload["summary"]["block_count"]),
        ("manual_review_block_count", payload["summary"]["manual_review_block_count"]),
        ("safety_table_count", payload["summary"]["safety_table_count"]),
        ("navigation_ui_count", payload["summary"]["navigation_ui_count"]),
        ("spec_table_count", payload["summary"]["spec_table_count"]),
    ]
    for row in rows:
        sheet.append(list(row))
    style_sheet(sheet, [28, 90])


def write_content_review_sheet(sheet, payload: dict[str, Any]) -> None:
    headers = [
        "block_order",
        "section_heading",
        "block_type",
        "manual_review_required",
        "text",
        "sentence_count",
        "crop_path",
        "page",
        "cell_order",
    ]
    sheet.append(headers)
    for block in payload["blocks"]:
        sheet.append(
            [
                block["block_order"],
                block["section_heading"],
                block["block_type"],
                block["manual_review_required"],
                block["text"],
                len(block["sentences"]),
                block.get("crop_path") or "",
                block["source"]["page_number"],
                block["source"]["cell_order"],
            ]
        )
    style_sheet(sheet, [12, 42, 22, 22, 120, 15, 80, 8, 12])


def write_safety_tables_sheet(sheet, payload: dict[str, Any]) -> None:
    headers = ["block_order", "section_heading", "row_order", "symbol_key", "label", "description"]
    sheet.append(headers)
    for block in payload["blocks"]:
        if block["block_type"] != "safety_symbol_table":
            continue
        for row in block.get("rows", []):
            sheet.append(
                [
                    block["block_order"],
                    block["section_heading"],
                    row["row_order"],
                    row["symbol_key"],
                    row["label"],
                    row["description"],
                ]
            )
    style_sheet(sheet, [12, 42, 10, 24, 42, 120])


def write_navigation_ui_sheet(sheet, payload: dict[str, Any]) -> None:
    headers = ["block_order", "section_heading", "text", "crop_path", "page", "cell_order"]
    sheet.append(headers)
    for block in payload["blocks"]:
        if block["block_type"] != "navigation_ui":
            continue
        sheet.append(
            [
                block["block_order"],
                block["section_heading"],
                block["text"],
                block.get("crop_path") or "",
                block["source"]["page_number"],
                block["source"]["cell_order"],
            ]
        )
    style_sheet(sheet, [12, 42, 120, 80, 8, 12])


def write_spec_tables_sheet(sheet, payload: dict[str, Any]) -> None:
    headers = ["block_order", "section_heading", "row_order", "field", "value", "crop_path"]
    sheet.append(headers)
    for block in payload["blocks"]:
        if block["block_type"] != "spec_table":
            continue
        for row in block.get("rows", []):
            sheet.append(
                [
                    block["block_order"],
                    block["section_heading"],
                    row["row_order"],
                    row["field"],
                    row["value"],
                    block.get("crop_path") or "",
                ]
            )
    style_sheet(sheet, [12, 42, 10, 30, 120, 80])


def style_sheet(sheet, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def render_content_cell_crops(pdf_path: Path, payload: dict[str, Any], paths: ContentPocPaths) -> None:
    paths.crop_dir.mkdir(parents=True, exist_ok=True)
    zoom = 2
    matrix = fitz.Matrix(zoom, zoom)
    crop_cache: dict[tuple[int, int], str] = {}
    with fitz.open(pdf_path) as document:
        for block in payload["blocks"]:
            source = block["source"]
            key = (source["page_number"], source["cell_order"])
            if key not in crop_cache:
                page = document[source["page_number"] - 1]
                rect = cell_rect(page.rect, source["row"], source["column"], rows=2, columns=8)
                crop_path = paths.crop_dir / (
                    f"page_{source['page_number']:03d}_cell_{source['cell_order']:02d}.png"
                )
                page.get_pixmap(matrix=matrix, clip=rect, alpha=False).save(crop_path)
                crop_cache[key] = str(crop_path)
            if block["image_crop_required"]:
                block["crop_path"] = crop_cache[key]


def cell_rect(page_rect: fitz.Rect, row: int, column: int, rows: int, columns: int) -> fitz.Rect:
    cell_width = page_rect.width / columns
    cell_height = page_rect.height / rows
    return fitz.Rect(
        page_rect.x0 + (column - 1) * cell_width,
        page_rect.y0 + (row - 1) * cell_height,
        page_rect.x0 + column * cell_width,
        page_rect.y0 + row * cell_height,
    )


def run_content_poc(
    pdf_path: str | Path,
    output_dir: str | Path,
    language: str = "ENG",
    mapping_path: str | Path = DEFAULT_MAPPING_PATH,
) -> dict[str, Path]:
    pdf_path = Path(pdf_path)
    repository = PdfProfileRepository(mapping_path)
    lookup = ProfileLookupService(repository).lookup_for_file(str(pdf_path))
    structure = analyze_pdf_structure(str(pdf_path))
    extracted = extract_review_text(
        lookup.parsed_filename,
        lookup.profile,
        structure,
        language=language,
    )
    payload = build_content_poc_payload(extracted)
    paths = make_output_paths(Path(output_dir), extracted["source_token"], language)
    render_content_cell_crops(pdf_path, payload, paths)
    return write_content_poc_outputs(payload, paths)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Generate a content-only POC review package.")
    parser.add_argument("pdf_path", help="Path to the source PDF.")
    parser.add_argument("--language", default="ENG", help="Language code to extract.")
    parser.add_argument(
        "--output-dir",
        default="outputs/content_poc",
        help="Root output directory. A source/language subfolder is created inside it.",
    )
    parser.add_argument(
        "--mapping",
        default=str(DEFAULT_MAPPING_PATH),
        help=f"Path to profile mapping JSON. Default: {DEFAULT_MAPPING_PATH}",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    written = run_content_poc(
        args.pdf_path,
        args.output_dir,
        language=args.language,
        mapping_path=args.mapping,
    )
    for label, path in written.items():
        print(f"{label}: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

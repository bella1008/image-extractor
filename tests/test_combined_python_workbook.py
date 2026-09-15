"""The deployable writer preserves the reviewer contract without a Node runtime."""
import hashlib
import importlib
import json

from openpyxl import load_workbook
import pytest

from tests.test_combined_workbook_backend import view
from tests.test_combined_review_model import combined_inputs
from tests.test_item_review_service import checked_bundle


def writer():
    assert importlib.util.find_spec('src.combined_review_workbook'), 'Python workbook writer missing'
    return importlib.import_module('src.combined_review_workbook')


def source(tmp_path, data):
    path = tmp_path / 'view.json'
    content = json.dumps(data, ensure_ascii=True).encode()
    path.write_bytes(content)
    return path, hashlib.sha256(content).hexdigest()


@pytest.mark.parametrize('empty', [False, True])
def test_real_python_workbook_preserves_every_cell_and_layout(tmp_path, empty):
    data = view()
    if empty:
        data['sheets'][-1]['rows'] = []
    else:
        data['sheets'][1]['rows'][0][4] = '#N/A'
        data['sheets'][2]['rows'][0][6] = '한글 café 中文\n* condition 😀'
    path, digest = source(tmp_path, data)
    writer().write_combined_workbook(path, tmp_path, digest)
    from src.item_review_excel import validate_item_workbook
    output = tmp_path / 'review_report.xlsx'
    validate_item_workbook(output.read_bytes(), data)
    wb = load_workbook(output)
    try:
        for ws, spec in zip(wb.worksheets, data['sheets'], strict=True):
            assert not ws.sheet_view.showGridLines
            assert ws['A1'].font.bold and ws['A1'].font.name == 'Arial'
            assert ws['A1'].fill.fgColor.rgb[-6:] == 'D9EAF7'
            assert ws.row_dimensions[1].height == 32
            assert ws['A1'].alignment.wrap_text
            assert all(d.height <= 409 for d in ws.row_dimensions.values())
            for col, width in enumerate(spec['widths'], 1):
                from openpyxl.utils import get_column_letter
                assert ws.column_dimensions[get_column_letter(col)].width == width
            if spec['rows'] and '검토 판정' in spec['headers']:
                col = spec['headers'].index('검토 판정') + 1
                assert ws.cell(2, col).fill.fgColor.rgb[-6:] == 'FFF4CE'
    finally:
        wb.close()
    before = output.read_bytes()
    with pytest.raises(FileExistsError):
        writer().write_combined_workbook(path, tmp_path, digest)
    assert output.read_bytes() == before
    assert not list(tmp_path.glob('*.pending.xlsx'))


@pytest.mark.parametrize('fault', ['schema', 'state', 'headers', 'sheets', 'summary', 'width',
                                  'freeze', 'height', 'long_cell', 'control', 'float', 'bool', 'row'])
def test_invalid_view_never_creates_workbook(tmp_path, fault):
    data = view()
    sheet = data['sheets'][1]
    if fault == 'schema': data['schema_version'] = 'combined-review-excel-view/99'
    elif fault == 'state': data['decision_status'] = 'pass'
    elif fault == 'headers': sheet['headers'][0] = 'unknown'
    elif fault == 'sheets': data['sheets'].reverse()
    elif fault == 'summary': data['sheets'][0]['rows'].pop()
    elif fault == 'width': sheet['widths'][0] = 0
    elif fault == 'freeze': sheet['freeze'] = 'Z500'
    elif fault == 'height': sheet['rows'][0][3] = 'line\n' * 32
    elif fault == 'long_cell': sheet['rows'][0][3] = '😀' * 16384
    elif fault == 'control': sheet['rows'][0][3] = 'bad\x00text'
    elif fault == 'float': sheet['rows'][0][3] = 1.5
    elif fault == 'bool': sheet['rows'][0][3] = True
    elif fault == 'row': sheet['rows'][0].pop()
    path, digest = source(tmp_path, data)
    with pytest.raises(ValueError):
        writer().write_combined_workbook(path, tmp_path, digest)
    assert not (tmp_path / 'review_report.xlsx').exists()


def test_changed_input_is_rejected_before_publication(tmp_path, monkeypatch):
    path, digest = source(tmp_path, view())
    backend = writer()
    with pytest.raises(ValueError, match='hash|changed'):
        backend.write_combined_workbook(path, tmp_path, '0' * 64)
    save = backend.Workbook.save
    def changed(self, target):
        save(self, target)
        path.write_bytes(path.read_bytes() + b' ')
    monkeypatch.setattr(backend.Workbook, 'save', changed)
    with pytest.raises(ValueError, match='hash|changed'):
        backend.write_combined_workbook(path, tmp_path, digest)
    assert not (tmp_path / 'review_report.xlsx').exists()
    assert not list(tmp_path.glob('*.pending.xlsx'))


def test_combined_service_exports_without_node(combined_inputs, tmp_path, monkeypatch):
    from src import combined_review_service as service
    def forbidden(*args, **kwargs):
        pytest.fail('Python export must not start Node or a subprocess')
    monkeypatch.setattr(service.subprocess, 'run', forbidden)
    output = tmp_path / 'python-result'
    service.export_combined_review(tmp_path / 'checklist', tmp_path / 'run', output)
    screen = service.read_completed_combined_review(output)
    assert screen['report']['summary']['parent_check_count'] == 59
    assert screen['report']['summary']['child_item_count'] == 14
    assert not list(output.glob('*.html'))


def test_cli_defaults_to_python_and_rejects_half_node_configuration(monkeypatch, tmp_path):
    from scripts import run_combined_review_v2 as command
    seen = []
    monkeypatch.setenv('ITEM_REVIEW_NODE', 'obsolete-developer-path')
    monkeypatch.setenv('ITEM_REVIEW_NODE_MODULES', 'obsolete-developer-path')
    monkeypatch.setattr(command, 'run_combined_review', seen.append)
    assert command.main(['test.pdf', '--output', str(tmp_path)]) == 0
    assert seen[0].node_executable is None and seen[0].node_modules is None
    with pytest.raises(SystemExit) as error:
        command.main(['test.pdf', '--output', str(tmp_path), '--node', 'node'])
    assert error.value.code == 2

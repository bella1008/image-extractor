"""Separate, non-active item authoring draft with externally pinned source data."""
from io import BytesIO
import hashlib
import json
import os
from pathlib import Path
import tempfile

from src.checklist_item_proposal import ItemSlice, build_item_proposal
from src.checklist_migration import DRAFT_HEADERS

EDITABLE_FIELDS = ('proposed_model_rule', 'proposal_evidence', 'reviewer_note')
ITEM_HEADERS = ('source_check_id', 'item_key', 'required_text', 'model_applicability', 'condition_state',
                *EDITABLE_FIELDS, 'start', 'end', 'common_id', 'language', 'scope', 'exclude_scope',
                'source_marker', 'evidence_refs', 'condition_refs')
EVIDENCE_HEADERS = ('evidence_id', 'kind', 'item_key', 'text', 'source_node_ids', 'structure_types',
                    'pdf_pages', 'details_json')
SHEET_HEADERS = {'Items': ITEM_HEADERS, 'Parent': ('field', 'value'), 'Source Evidence': EVIDENCE_HEADERS,
                 'Source': ('field', 'value'), 'Guide': ('field', 'value')}


def _require(condition, message):
    if not condition:
        raise ValueError(message)


def _json(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'))


def build_master_seed(proposal: dict) -> dict:
    """Use a freshly checked proposal; this function does not establish provenance."""
    _require(proposal['activation_status'] == 'proposal_only' and
             proposal['decision_status'] == 'not_evaluated', 'active proposal forbidden')
    parent = proposal['source_rule']
    rebuilt = build_item_proposal(parent, tuple(ItemSlice(i['item_key'], i['start'], i['end'])
                                               for i in proposal['items']))
    _require(parent['check_id'] == 'CHK-002-ZC-ENG' and len(rebuilt['items']) == 14, 'CHK-002 pilot required')
    _require(proposal['source_segments'] == rebuilt['source_segments'], 'source segments differ')
    items, evidence = [], []
    for original, expected in zip(proposal['items'], rebuilt['items'], strict=True):
        _require(all(original[k] == v for k, v in expected.items()), 'source item differs')
        row = {k: str(v) for k, v in expected.items() if k in ITEM_HEADERS}
        row.update(model_applicability='unknown', condition_state='unverified',
                   source_marker=original['source_marker'], **dict.fromkeys(EDITABLE_FIELDS, ''))
        for kind, input_key, ref_key in (('item', 'candidates', 'evidence_refs'),
                                         ('condition_candidate', 'condition_candidates', 'condition_refs')):
            refs = []
            for index, window in enumerate(original[input_key], 1):
                ref = f'{original["item_key"]}:{kind}:{index}'
                refs.append(ref)
                pages = sorted({e['page_index'] + 1 for p in window['parts']
                                for e in p.get('evidence', []) if e.get('page_index') is not None})
                evidence.append([ref, kind, row['item_key'], window['text'],
                                 ', '.join(window['owner_ids']), ', '.join(window['structure_types']),
                                 ', '.join(map(str, pages)), _json(window)])
            row[ref_key] = _json(refs)
        items.append([row[k] for k in ITEM_HEADERS])
    source = dict(proposal['source'])
    source.update(source_segments_json=_json(proposal['source_segments']),
                  scope_observation_json=_json(proposal['scope_observation']),
                  comparison_policy=proposal['comparison_policy'])
    _require(all(isinstance(v, str) for v in source.values()), 'source values must be text')
    rows = {'Items': items, 'Parent': [[k, parent[k]] for k in DRAFT_HEADERS],
            'Source Evidence': evidence, 'Source': [[k, v] for k, v in source.items()],
            'Guide': [
                ['purpose', 'CHK-002 항목별 원장 작성 초안입니다. 기존 547행 DB와 자동 검토에는 연결되지 않습니다.'],
                ['editable_fields', 'Items의 proposed_model_rule, proposal_evidence, reviewer_note만 수정하세요.'],
                ['proposed_model_rule', '모델 적용 조건 제안문입니다. 실행 규칙이나 승인으로 사용하지 않습니다.'],
                ['proposal_evidence', '제안의 근거를 적으세요. source evidence와 별개인 검토자 입력입니다.'],
                ['reviewer_note', '검토 의견을 적으세요. 빈칸도 허용합니다.'],
                ['unknown / unverified', '모델 적용 여부는 미확정(unknown), 조건 관계는 미검증(unverified)입니다.'],
                ['source_fields', '원문·별표·수량·괄호·범위·출처는 변경할 수 없습니다. 항목 행의 순서만 바꿀 수 있습니다.'],
                ['details_json', 'Source Evidence의 details_json은 노드·문자 구간·페이지·원문을 포함한 완전한 기계 감사 근거입니다.'],
                ['draft_only', '부모의 과거 approved 값은 원본 이력입니다. 이 항목 초안의 승인이나 Pass 판정이 아닙니다.'],
                ['export', '원장 Excel과 별도 보관한 seed의 SHA-256을 검증하여 새 draft JSON으로 내보냅니다.'],
            ]}
    return {'schema_version': 'checklist-item-master-seed/1', 'activation_status': 'draft_only',
            'sheets': [{'name': name, 'headers': list(headers), 'rows': rows[name]}
                       for name, headers in SHEET_HEADERS.items()]}


def publish_json(output: Path, payload: dict, verify_inputs=lambda: None) -> None:
    """Publish one complete file without replacing an existing destination."""
    output = Path(output)
    if output.exists():
        raise FileExistsError(f'use a fresh output file: {output}')
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=output.name + '.', suffix='.pending', dir=output.parent)
    pending = Path(name)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8', newline='\n') as stream:
            json.dump(payload, stream, ensure_ascii=True, indent=2)
            stream.write('\n')
            stream.flush()
            os.fsync(stream.fileno())
        verify_inputs()
        os.link(pending, output)
    finally:
        pending.unlink(missing_ok=True)


def _read_sheets(excel_bytes, seed):
    from openpyxl import load_workbook
    _require(seed.get('schema_version') == 'checklist-item-master-seed/1' and
             seed.get('activation_status') == 'draft_only', 'unsupported or active seed')
    expected_sheets = seed['sheets']
    _require([s['name'] for s in expected_sheets] == list(SHEET_HEADERS), 'seed sheets differ')
    workbook = load_workbook(BytesIO(excel_bytes), read_only=True, data_only=False)
    result = {}
    try:
        _require(workbook.sheetnames == list(SHEET_HEADERS), 'workbook sheets differ')
        for definition in expected_sheets:
            name, headers, expected = definition['name'], definition['headers'], definition['rows']
            _require(headers == list(SHEET_HEADERS[name]), f'seed headers differ: {name}')
            actual = []
            # Ignore producer dimension hints; inspect every represented cell.
            workbook[name].reset_dimensions()
            for cells in workbook[name].iter_rows():
                values = []
                for cell in cells:
                    _require(cell.data_type not in ('f', 'e'), f'formula/error cell: {name}')
                    _require(cell.value is None or isinstance(cell.value, str), f'non-text cell: {name}')
                    values.append('' if cell.value is None else cell.value)
                # Omitted trailing empty cells are legitimate OOXML blanks.
                _require(len(values) <= len(headers), f'extra column: {name}')
                actual.append(values + [''] * (len(headers) - len(values)))
            _require(actual and actual[0] == headers, f'headers differ: {name}')
            actual = actual[1:]
            _require(len(actual) == len(expected), f'row count differs: {name}')
            if name == 'Items':
                key_index = headers.index('item_key')
                keyed = {r[key_index]: r for r in actual}
                _require(len(keyed) == len(actual) and set(keyed) == {r[key_index] for r in expected},
                         'missing, extra or duplicate item key')
                actual = [keyed[r[key_index]] for r in expected]
            for row, baseline in zip(actual, expected, strict=True):
                _require(len(baseline) == len(headers) and all(isinstance(v, str) for v in baseline),
                         f'invalid seed row: {name}')
                for column, value, original in zip(headers, row, baseline, strict=True):
                    if name != 'Items' or column not in EDITABLE_FIELDS:
                        _require(value == original, f'immutable source changed: {name}.{column}')
            result[name] = [dict(zip(headers, row, strict=True)) for row in actual]
    finally:
        workbook.close()
    return result


def export_item_master(excel: Path, seed: Path, output: Path, *, expected_seed_sha256: str) -> dict:
    """Export reviewer text; pin the seed digest separately from editable files."""
    excel, seed = Path(excel), Path(seed)
    if Path(output).exists():
        raise FileExistsError(f'use a fresh output file: {output}')
    before = {path: path.read_bytes() for path in (excel, seed)}
    digest = hashlib.sha256(before[seed]).hexdigest()
    _require(digest == expected_seed_sha256, 'seed SHA-256 differs from external expected hash')
    sheets = _read_sheets(before[excel], json.loads(before[seed]))
    parent = {r['field']: r['value'] for r in sheets['Parent']}
    source = {r['field']: r['value'] for r in sheets['Source']}
    items = sheets['Items']
    for item in items:
        for field in ('start', 'end'):
            item[field] = int(item[field])
        for field in ('evidence_refs', 'condition_refs'):
            item[field] = json.loads(item[field])
        _require(item['model_applicability'] == 'unknown' and item['condition_state'] == 'unverified',
                 'item applicability cannot be activated')
    result = {'schema_version': 'checklist-item-master-draft/1', 'activation_status': 'draft_only',
              'parent': parent, 'items': items, 'evidence': sheets['Source Evidence'], 'source': source,
              'provenance': {'seed_sha256': digest, 'excel_sha256': hashlib.sha256(before[excel]).hexdigest()}}
    def unchanged():
        _require(all(path.read_bytes() == data for path, data in before.items()), 'input changed during export')
    publish_json(output, result, unchanged)
    return result

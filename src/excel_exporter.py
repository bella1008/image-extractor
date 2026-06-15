from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTLE_FILL = PatternFill("solid", fgColor="D9EAF7")
KIND_FILLS = {
    "heading": PatternFill("solid", fgColor="D9EAF7"),
    "body": PatternFill("solid", fgColor="FFFFFF"),
    "bullet": PatternFill("solid", fgColor="E2F0D9"),
    "warning": PatternFill("solid", fgColor="FCE4D6"),
    "navigation": PatternFill("solid", fgColor="E4DFEC"),
    "navigation_ui": PatternFill("solid", fgColor="E4DFEC"),
    "list_item": PatternFill("solid", fgColor="FFF2CC"),
    "ui_label": PatternFill("solid", fgColor="DDEBF7"),
    "safety_symbol_table": PatternFill("solid", fgColor="F8CBAD"),
}


def export_review_text_xlsx(json_path: str | Path, output_path: str | Path) -> Path:
    json_path = Path(json_path)
    output_path = Path(output_path)
    data = json.loads(json_path.read_text(encoding="utf-8"))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    write_summary(summary, data)

    write_review_items(workbook.create_sheet("Review Items"), data)
    write_cover_regions(workbook.create_sheet("Cover Regions"), data)
    write_section_blocks(workbook.create_sheet("Section Blocks"), data)
    write_safety_tables(workbook.create_sheet("Safety Tables"), data)
    write_normalized_sentences(workbook.create_sheet("Normalized Sentences"), data)
    write_cells(workbook.create_sheet("Cells"), data)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def write_summary(sheet, data: dict[str, Any]) -> None:
    rows = [
        ("File Name", data["file_name"]),
        ("Source Token", data["source_token"]),
        ("Region", data["region"]),
        ("Buyer Codes", ";".join(data["buyer_codes"])),
        ("Language", data["language"]),
        ("Language Variant", data["language_variant"] or ""),
        ("Doc Type", data["doc_type"]),
        ("Detected Doc Type", data["detected_doc_type"]),
        ("Page Count", data["page_count"]),
        ("Extraction Method", data["extraction"]["method"]),
        ("OCR Applied", data["extraction"]["ocr_applied"]),
        ("OCR Recommended", data["extraction"]["ocr_recommended"]),
        ("Total Chars", data["extraction"]["total_chars"]),
        ("Total Sentences", data["extraction"]["total_sentences"]),
        ("Total Cells", data["extraction"]["total_cells"]),
        ("Empty Cells", data["extraction"]["empty_cells"]),
        ("Sections", len(data.get("sections", []))),
        ("Cover Pages", len(data.get("cover_pages", []))),
    ]
    sheet.append(["Field", "Value"])
    for row in rows:
        sheet.append(list(row))
    style_header(sheet, 1, 2)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 90
    sheet.freeze_panes = "A2"


def write_review_items(sheet, data: dict[str, Any]) -> None:
    headers = [
        "page",
        "review_order",
        "row",
        "column",
        "cell_role",
        "cell_order",
        "item_order",
        "kind",
        "text",
    ]
    sheet.append(headers)
    for page in data["pages"]:
        for cell in page["cells"]:
            for item in cell["review_items"]:
                sheet.append(
                    [
                        page["page_number"],
                        page["review_order"],
                        cell["row"],
                        cell["column"],
                        cell["role"],
                        cell["reading_order"],
                        item["item_order"],
                        item["kind"],
                        item["text"],
                    ]
                )
    style_table(sheet, len(headers), kind_col=8)
    set_widths(sheet, [8, 13, 8, 8, 16, 11, 11, 14, 100])


def write_cover_regions(sheet, data: dict[str, Any]) -> None:
    headers = [
        "page",
        "schema_id",
        "schema_status",
        "missing_required_regions",
        "region_order",
        "region_type",
        "heading",
        "cell_order",
        "bbox",
        "bold",
        "max_font_size",
        "text",
        "markdown",
    ]
    sheet.append(headers)
    for cover in data.get("cover_pages", []):
        schema = cover.get("schema", {})
        validation = cover.get("schema_validation", {})
        for region in cover["regions"]:
            style = region.get("style_summary", {})
            sheet.append(
                [
                    cover["page_number"],
                    schema.get("schema_id", ""),
                    validation.get("status", ""),
                    ";".join(validation.get("missing_required_regions", [])),
                    region["region_order"],
                    region["region_type"],
                    region.get("heading") or "",
                    region["source"]["cell_order"],
                    ",".join(str(value) for value in region["bbox"]),
                    style.get("bold_candidate", False),
                    style.get("max_font_size", 0),
                    region["text"],
                    region["markdown"],
                ]
            )
    style_table(sheet, len(headers), kind_col=6)
    set_widths(sheet, [8, 18, 16, 30, 13, 20, 32, 11, 24, 8, 13, 100, 100])


def write_normalized_sentences(sheet, data: dict[str, Any]) -> None:
    headers = [
        "page",
        "review_order",
        "row",
        "column",
        "cell_role",
        "cell_order",
        "sentence_order",
        "text",
    ]
    sheet.append(headers)
    for page in data["pages"]:
        for cell in page["cells"]:
            for sentence in cell["normalized_sentences"]:
                sheet.append(
                    [
                        page["page_number"],
                        page["review_order"],
                        cell["row"],
                        cell["column"],
                        cell["role"],
                        cell["reading_order"],
                        sentence["sentence_order"],
                        sentence["text"],
                    ]
                )
    style_table(sheet, len(headers))
    set_widths(sheet, [8, 13, 8, 8, 16, 11, 14, 120])


def write_section_blocks(sheet, data: dict[str, Any]) -> None:
    headers = [
        "section_order",
        "heading",
        "block_order",
        "block_type",
        "page",
        "row",
        "column",
        "cell_order",
        "bold",
        "max_font_size",
        "text",
    ]
    sheet.append(headers)
    for section in data.get("sections", []):
        for block in section["blocks"]:
            source = block["source"]
            style = block.get("style_summary", {})
            sheet.append(
                [
                    section["section_order"],
                    section["heading"],
                    block["block_order"],
                    block["block_type"],
                    source["page_number"],
                    source["row"],
                    source["column"],
                    source["cell_order"],
                    style.get("bold_candidate", False),
                    style.get("max_font_size", 0),
                    block["text"],
                ]
            )
    style_table(sheet, len(headers), kind_col=4)
    set_widths(sheet, [14, 42, 12, 20, 8, 8, 8, 11, 8, 13, 120])


def write_safety_tables(sheet, data: dict[str, Any]) -> None:
    headers = [
        "section_order",
        "heading",
        "block_order",
        "row_order",
        "symbol_key",
        "label",
        "lines_text",
        "image_crop_required",
        "ocr_recommended",
        "manual_review_required",
    ]
    sheet.append(headers)
    for section in data.get("sections", []):
        for block in section["blocks"]:
            if block["block_type"] != "safety_symbol_table":
                continue
            for row in block.get("rows", []):
                sheet.append(
                    [
                        section["section_order"],
                        section["heading"],
                        block["block_order"],
                        row["row_order"],
                        row["symbol_key"],
                        row["label"],
                        "\n".join(row.get("lines", [])),
                        block.get("image_crop_required", False),
                        block.get("ocr_recommended", False),
                        block.get("manual_review_required", False),
                    ]
                )
    style_table(sheet, len(headers))
    set_widths(sheet, [14, 42, 12, 10, 22, 38, 120, 18, 16, 22])


def write_cells(sheet, data: dict[str, Any]) -> None:
    headers = [
        "page",
        "review_order",
        "row",
        "column",
        "cell_role",
        "cell_order",
        "direction",
        "char_count",
        "span_count",
        "bold_span_count",
        "max_font_size",
        "normalized_text",
        "raw_text",
    ]
    sheet.append(headers)
    for page in data["pages"]:
        for cell in page["cells"]:
            sheet.append(
                [
                    page["page_number"],
                    page["review_order"],
                    cell["row"],
                    cell["column"],
                    cell["role"],
                    cell["reading_order"],
                    cell["reading_direction"],
                    cell["char_count"],
                    cell.get("style_summary", {}).get("span_count", 0),
                    cell.get("style_summary", {}).get("bold_span_count", 0),
                    cell.get("style_summary", {}).get("max_font_size", 0),
                    cell["normalized_text"],
                    cell["text"],
                ]
            )
    style_table(sheet, len(headers))
    set_widths(sheet, [8, 13, 8, 8, 16, 11, 11, 11, 11, 15, 13, 120, 120])


def style_table(sheet, header_count: int, kind_col: int | None = None) -> None:
    style_header(sheet, 1, header_count)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if kind_col:
            kind = row[kind_col - 1].value
            fill = KIND_FILLS.get(kind)
            if fill:
                row[kind_col - 1].fill = fill


def style_header(sheet, row_number: int, header_count: int) -> None:
    for col in range(1, header_count + 1):
        cell = sheet.cell(row=row_number, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def set_widths(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

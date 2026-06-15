import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.cover_poc import (
    build_cover_poc_payload,
    build_cover_review_markdown,
    make_output_paths,
    safe_region_slug,
    write_cover_poc_outputs,
)


class CoverPocTests(unittest.TestCase):
    def test_make_output_paths_uses_source_token_and_language(self):
        paths = make_output_paths(Path("outputs/cover_poc"), "ZC_L02", "ENG")

        self.assertEqual(paths.root, Path("outputs/cover_poc/ZC_L02_ENG"))
        self.assertEqual(paths.json_path.name, "cover_poc.json")
        self.assertEqual(paths.markdown_path.name, "cover_poc.md")
        self.assertEqual(paths.xlsx_path.name, "cover_regions.xlsx")
        self.assertEqual(paths.page_image_path.name, "cover_page_001.png")
        self.assertEqual(paths.region_dir.name, "regions")

    def test_safe_region_slug_keeps_order_type_and_short_text(self):
        slug = safe_region_slug(
            {
                "region_order": 9,
                "region_type": "contact_table",
                "heading": "Contact table",
                "text": "Samsung Service Center / Website",
            }
        )

        self.assertEqual(slug, "region_09_contact_table_contact_table")

    def test_build_cover_poc_payload_keeps_cover_schema_and_regions(self):
        extracted = minimal_extracted_payload()

        payload = build_cover_poc_payload(extracted)

        self.assertEqual(payload["file_name"], "sample.pdf")
        self.assertEqual(payload["source_token"], "ZC_L02")
        self.assertEqual(payload["language"], "ENG")
        self.assertEqual(payload["cover"]["schema"]["schema_id"], "ZC_A2_COVER")
        self.assertEqual(payload["cover"]["schema_validation"]["status"], "PASS")
        self.assertEqual(payload["summary"]["region_count"], 3)
        self.assertEqual(payload["summary"]["manual_review_region_count"], 2)
        self.assertEqual(payload["regions"][2]["region_type"], "contact_table")

    def test_build_cover_review_markdown_contains_table_and_review_flags(self):
        payload = build_cover_poc_payload(minimal_extracted_payload())

        markdown = build_cover_review_markdown(payload)

        self.assertIn("# Cover Page POC", markdown)
        self.assertIn("schema_id: `ZC_A2_COVER`", markdown)
        self.assertIn("manual_review_required: `true`", markdown)
        self.assertIn("| Samsung Service Center | Website |", markdown)

    def test_write_cover_poc_outputs_writes_json_markdown_and_xlsx(self):
        payload = build_cover_poc_payload(minimal_extracted_payload())

        with tempfile.TemporaryDirectory() as temp_dir:
            paths = make_output_paths(Path(temp_dir), "ZC_L02", "ENG")
            written = write_cover_poc_outputs(payload, paths)

            self.assertTrue(written["json_path"].exists())
            self.assertTrue(written["markdown_path"].exists())
            self.assertTrue(written["xlsx_path"].exists())
            saved = json.loads(written["json_path"].read_text(encoding="utf-8"))
            self.assertEqual(saved["cover"]["schema"]["schema_id"], "ZC_A2_COVER")

            workbook = load_workbook(written["xlsx_path"])
            self.assertEqual(
                workbook.sheetnames,
                ["Cover Summary", "Cover Review", "Cover Markdown", "Compare Fields"],
            )
            self.assertEqual(
                [cell.value for cell in workbook["Cover Review"][1]],
                [
                    "region_order",
                    "region_type",
                    "heading",
                    "manual_review_required",
                    "text",
                    "crop_path",
                ],
            )
            self.assertEqual(
                [cell.value for cell in workbook["Cover Markdown"][1]],
                ["region_order", "region_type", "markdown"],
            )
            self.assertEqual(
                [cell.value for cell in workbook["Compare Fields"][1]],
                ["region_order", "region_type", "field_schema", "field_name", "field_value", "label"],
            )


def minimal_extracted_payload():
    return {
        "file_name": "sample.pdf",
        "source_token": "ZC_L02",
        "region": "ZC",
        "buyer_codes": ["ZC"],
        "language": "ENG",
        "doc_type": "A2",
        "detected_doc_type": "A2",
        "cover_pages": [
            {
                "page_number": 1,
                "page_role": "cover",
                "schema": {"schema_id": "ZC_A2_COVER", "schema_version": 1},
                "schema_validation": {
                    "status": "PASS",
                    "missing_required_regions": [],
                    "present_regions": ["title", "contact_table"],
                },
                "markdown": "# Simple User Guide\n\n| Samsung Service Center | Website |\n| --- | --- |\n| 1-800 | www.samsung.com |",
                "regions": [
                    {
                        "region_order": 1,
                        "region_type": "title",
                        "heading": "Simple User Guide",
                        "text": "Simple User Guide",
                        "bbox": [50, 100, 300, 150],
                        "source": {"cell_order": 1},
                        "style_summary": {"bold_candidate": True, "max_font_size": 24},
                        "markdown": "# Simple User Guide",
                    },
                    {
                        "region_order": 2,
                        "region_type": "body",
                        "heading": None,
                        "text": "Thank you for purchasing this Samsung product.",
                        "bbox": [50, 150, 300, 180],
                        "source": {"cell_order": 1},
                        "style_summary": {"bold_candidate": False, "max_font_size": 8},
                        "markdown": "Thank you for purchasing this Samsung product.",
                    },
                    {
                        "region_order": 3,
                        "region_type": "contact_table",
                        "heading": "Contact table",
                        "text": "Samsung Service Center: 1-800 Website: www.samsung.com",
                        "bbox": [50, 250, 400, 320],
                        "source": {"cell_order": 2},
                        "style_summary": {},
                        "manual_review_required": True,
                        "markdown": "| Samsung Service Center | Website |\n| --- | --- |\n| 1-800 | www.samsung.com |",
                        "table": {
                            "table_type": "cover_contact_table",
                            "columns": ["Samsung Service Center", "Website"],
                            "rows": [
                                {
                                    "Samsung Service Center": "1-800",
                                    "Website": "www.samsung.com",
                                }
                            ],
                        },
                        "normalized_fields": {
                            "field_schema": "cover_contact_table.v1",
                            "service_center": ["1-800"],
                            "service_phone_numbers": ["1-800"],
                            "websites": [{"url": "www.samsung.com", "label": ""}],
                            "address": [],
                        },
                    },
                ],
            }
        ],
    }


if __name__ == "__main__":
    unittest.main()
